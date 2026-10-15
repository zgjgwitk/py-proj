"""
合并 dataFiles 目录下所有 Excel 文件到单个 CSV 文件
"""
import csv
import glob
import os
from datetime import datetime
from openpyxl import load_workbook

# 源目录与输出文件路径
SOURCE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "dataFiles")
OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))
# 输出文件名带时间戳，避免重复运行覆盖
TIMESTAMP = datetime.now().strftime("%Y%m%d_%H%M%S")
OUTPUT_FILE = os.path.join(OUTPUT_DIR, f"merged_{TIMESTAMP}.csv")
# 目标工作表名（源文件中的工作表名）
TARGET_SHEET_NAME = "日志结果"


def collect_excel_files(directory):
    """收集目录下所有 .xlsx 文件，按文件名排序保证顺序稳定"""
    pattern = os.path.join(directory, "*.xlsx")
    files = sorted(glob.glob(pattern))
    return files


def merge_exccles(source_files, output_file):
    """
    将多个 Excel 文件的所有数据合并到单个 CSV 文件
    首个文件的表头作为最终表头，后续文件跳过表头
    """
    total_rows = 0
    header_written = False
    output_file_handle = open(output_file, "w", newline="", encoding="utf-8-sig")
    csv_writer = None

    try:
        for src_path in source_files:
            file_name = os.path.basename(src_path)
            src_wb = load_workbook(src_path, read_only=True, data_only=True)
            src_ws = src_wb[TARGET_SHEET_NAME]

            file_data_rows = 0
            for idx, row in enumerate(src_ws.iter_rows(values_only=True)):
                # 首个文件的第一行作为表头写入，后续文件跳过表头
                if idx == 0:
                    if not header_written:
                        csv_writer = csv.writer(output_file_handle)
                        csv_writer.writerow(row)
                        header_written = True
                    continue
                csv_writer.writerow(row)
                file_data_rows += 1

            total_rows += file_data_rows
            src_wb.close()
            print(f"已合并: {file_name} (数据行数: {file_data_rows})")
    finally:
        output_file_handle.close()

    return total_rows


def main():
    print(f"源目录: {SOURCE_DIR}")
    excel_files = collect_excel_files(SOURCE_DIR)
    print(f"发现 {len(excel_files)} 个 Excel 文件")

    if not excel_files:
        print("目录中没有可合并的 Excel 文件")
        return

    print(f"输出文件: {OUTPUT_FILE}")
    data_rows = merge_exccles(excel_files, OUTPUT_FILE)
    print(f"合并完成,共写入 {data_rows} 条数据(不含表头)")
    print(f"输出文件: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
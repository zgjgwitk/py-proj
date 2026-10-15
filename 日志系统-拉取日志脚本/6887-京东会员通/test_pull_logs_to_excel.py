import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from pull_logs_to_excel import (
    HourLogExportController,
    HourWindow,
    HourWindowFactory,
    IncrementalExcelExporter,
    LogApiClient,
    LogEventParser,
    LogExportApplication,
    LogMessageRecord,
    LogResponseParser,
    QueryConfig,
    create_config_from_args,
)


class LogEventParserTests(unittest.TestCase):
    def test_parse_message_should_extract_old_cdpid_and_bind_vipid(self) -> None:
        source_message = (
            "历史绑定校验，BindId=6887|1|20240625|97826_dbjd,"
            "BindMobileMatch=True,VipMobileMatch=True,BindVipId=1891463"
        )

        record = LogEventParser().parse_message(source_message)

        self.assertEqual("97826_dbjd", record.old_cdp_id)
        self.assertEqual("1891463", record.bind_vip_id)
        self.assertEqual(source_message, record.source_message)

    def test_parse_message_should_raise_when_bindid_missing(self) -> None:
        with self.assertRaisesRegex(ValueError, "BindId"):
            LogEventParser().parse_message("历史绑定校验，BindVipId=1891463")


class LogResponseParserTests(unittest.TestCase):
    def test_parse_response_should_only_keep_message_field(self) -> None:
        payload = {
            "data": {
                "logEvents": [
                    {
                        "message": (
                            "历史绑定校验，BindId=6887|1|20240625|97826_dbjd,"
                            "BindMobileMatch=True,VipMobileMatch=True,BindVipId=1891463"
                        ),
                        "level": "INFO",
                    },
                    {
                        "message": (
                            "历史绑定校验，BindId=6887|1|20240625|12345_dbjd,"
                            "BindMobileMatch=True,VipMobileMatch=True,BindVipId=999"
                        )
                    },
                ]
            }
        }

        rows = LogResponseParser(LogEventParser()).parse(payload)

        self.assertEqual(2, len(rows))
        self.assertEqual("97826_dbjd", rows[0].old_cdp_id)
        self.assertEqual("999", rows[1].bind_vip_id)


class IncrementalExcelExporterTests(unittest.TestCase):
    def test_append_rows_should_preserve_existing_content_across_batches(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "result.xlsx"
            exporter = IncrementalExcelExporter(output_path)
            exporter.initialize()
            exporter.append_rows(
                [
                    LogMessageRecord(
                        old_cdp_id="97826_dbjd",
                        bind_vip_id="1891463",
                        source_message=(
                            "历史绑定校验，BindId=6887|1|20240625|97826_dbjd,"
                            "BindMobileMatch=True,VipMobileMatch=True,BindVipId=1891463"
                        ),
                    )
                ]
            )
            exporter.append_rows(
                [
                    LogMessageRecord(
                        old_cdp_id="12345_dbjd",
                        bind_vip_id="999",
                        source_message=(
                            "历史绑定校验，BindId=6887|1|20240625|12345_dbjd,"
                            "BindMobileMatch=True,VipMobileMatch=True,BindVipId=999"
                        ),
                    )
                ]
            )

            workbook = load_workbook(output_path)
            try:
                sheet = workbook.active
                self.assertEqual("日志结果", sheet.title)
                self.assertEqual(["Old_CdpId", "BindVipId", "source_message"], [cell.value for cell in sheet[1]])
                self.assertEqual("97826_dbjd", sheet["A2"].value)
                self.assertEqual("1891463", sheet["B2"].value)
                self.assertEqual("12345_dbjd", sheet["A3"].value)
                self.assertEqual("999", sheet["B3"].value)
            finally:
                workbook.close()


class LogExportApplicationTests(unittest.TestCase):
    def test_run_should_flush_rows_by_batch_size(self) -> None:
        class FakeClient:
            def fetch_pages(self):
                yield [
                    LogMessageRecord("a_dbjd", "1", "message-1"),
                    LogMessageRecord("b_dbjd", "2", "message-2"),
                ]
                yield [
                    LogMessageRecord("c_dbjd", "3", "message-3"),
                ]

            def close(self) -> None:
                return None

        class FakeExporter:
            def __init__(self) -> None:
                self.initialized = False
                self.batches: list[list[str]] = []

            def initialize(self) -> None:
                self.initialized = True

            def append_rows(self, rows: list[LogMessageRecord]) -> None:
                self.batches.append([row.source_message for row in rows])

        config = QueryConfig(
            token="custom-token",
            output_path=Path("result.xlsx"),
            page_size=2,
            excel_batch_size=2,
        )
        exporter = FakeExporter()

        LogExportApplication(config, client=FakeClient(), exporter=exporter).run()

        self.assertTrue(exporter.initialized)
        self.assertEqual([["message-1", "message-2"], ["message-3"]], exporter.batches)


class ConfigTests(unittest.TestCase):
    def test_create_config_from_args_should_override_defaults(self) -> None:
        config = create_config_from_args(
            [
                "--token",
                "custom-token",
                "--cookie",
                "foo=bar",
                "--output",
                "custom.xlsx",
                "--excel-batch-size",
                "200",
            ]
        )

        self.assertEqual("custom-token", config.token)
        self.assertEqual("foo=bar", config.cookie)
        self.assertEqual("custom.xlsx", config.output_path.name)
        self.assertEqual(200, config.excel_batch_size)


class HourWindowFactoryTests(unittest.TestCase):
    @staticmethod
    def _format_hour_label(timestamp_ms: int) -> str:
        return datetime.fromtimestamp(timestamp_ms / 1000).strftime("%Y%m%d_%H")

    def test_split_should_create_one_window_per_hour(self) -> None:
        start_ms = 1_789_387_200_000
        end_ms = 1_789_390_799_000
        expected_label = self._format_hour_label(start_ms)

        windows = HourWindowFactory.split(start_date_ms=start_ms, end_date_ms=end_ms)

        self.assertEqual(1, len(windows))
        self.assertEqual(start_ms, windows[0].start_date_ms)
        self.assertEqual(end_ms, windows[0].end_date_ms)
        self.assertEqual(expected_label, windows[0].hour_label)

    def test_split_should_cover_multiple_hours(self) -> None:
        start_ms = 1_789_387_200_000
        end_ms = 1_789_397_999_000

        windows = HourWindowFactory.split(start_date_ms=start_ms, end_date_ms=end_ms)

        self.assertEqual(3, len(windows))
        self.assertEqual(start_ms, windows[0].start_date_ms)
        self.assertEqual(self._format_hour_label(start_ms), windows[0].hour_label)
        self.assertEqual(1_789_390_800_000, windows[1].start_date_ms)
        self.assertEqual(self._format_hour_label(1_789_390_800_000), windows[1].hour_label)
        self.assertEqual(1_789_394_400_000, windows[2].start_date_ms)
        self.assertEqual(self._format_hour_label(1_789_394_400_000), windows[2].hour_label)

        self.assertEqual(1_789_390_799_999, windows[0].end_date_ms)
        self.assertEqual(1_789_394_399_999, windows[1].end_date_ms)
        self.assertEqual(end_ms, windows[2].end_date_ms)


class HourLogExportControllerTests(unittest.TestCase):
    @staticmethod
    def _format_hour_label(timestamp_ms: int) -> str:
        return datetime.fromtimestamp(timestamp_ms / 1000).strftime("%Y%m%d_%H")

    def test_run_should_create_one_excel_per_hour_window(self) -> None:
        first_window_start = 1_789_387_200_000
        first_window_end = 1_789_390_799_999
        second_window_start = 1_789_390_800_000
        second_window_end = 1_789_394_399_999
        first_label = self._format_hour_label(first_window_start)
        second_label = self._format_hour_label(second_window_start)
        windows = [
            HourWindow(first_window_start, first_window_end, first_label),
            HourWindow(second_window_start, second_window_end, second_label),
        ]

        class FakeWindowFactory:
            @staticmethod
            def split(start_date_ms: int, end_date_ms: int) -> list[HourWindow]:
                return list(windows)

        class FakeApplication:
            def __init__(self) -> None:
                self.calls: list[tuple[int, int, Path]] = []

            def __call__(self, config: QueryConfig) -> LogExportApplication:
                captured = self.calls

                class InnerApp:
                    def run(inner_self) -> Path:
                        output_path = config.output_path
                        output_path.parent.mkdir(parents=True, exist_ok=True)
                        workbook = Workbook()
                        sheet = workbook.active
                        sheet.title = "日志结果"
                        sheet.append(("Old_CdpId", "BindVipId", "source_message"))
                        sheet.append(
                            (f"row-{config.start_date_ms}", "1", f"msg-{config.start_date_ms}")
                        )
                        workbook.save(output_path)
                        workbook.close()
                        captured.append((config.start_date_ms, config.end_date_ms, output_path))
                        return output_path

                return InnerApp()

        with tempfile.TemporaryDirectory() as temp_dir:
            base_output = Path(temp_dir) / "6887-京东会员通-日志结果.xlsx"
            config = QueryConfig(
                token="custom-token",
                start_date_ms=windows[0].start_date_ms,
                end_date_ms=windows[-1].end_date_ms,
                output_path=base_output,
            )
            controller = HourLogExportController(
                base_config=config,
                window_factory=FakeWindowFactory,
                application_factory=FakeApplication(),
            )

            output_paths = controller.run()

            self.assertEqual(2, len(output_paths))
            self.assertEqual(
                f"6887-京东会员通-日志结果_{first_label}.xlsx", output_paths[0].name
            )
            self.assertEqual(
                f"6887-京东会员通-日志结果_{second_label}.xlsx", output_paths[1].name
            )
            for path in output_paths:
                self.assertTrue(path.exists())


class LogApiClientCompatibilityTests(unittest.TestCase):
    @staticmethod
    def _build_message(old_cdp_id: str, bind_vip_id: str) -> str:
        return (
            f"历史绑定校验，BindId=6887|1|20240625|{old_cdp_id},"
            f"BindMobileMatch=True,VipMobileMatch=True,BindVipId={bind_vip_id}"
        )

    def test_fetch_pages_should_treat_missing_log_events_as_empty_and_stop(self) -> None:
        from unittest.mock import MagicMock

        def post_side_effect(url: str, json: dict[str, object], timeout: int) -> MagicMock:
            response = MagicMock()
            response.raise_for_status.return_value = None
            response.json.return_value = {"status": True, "data": {"total": 0}}
            response.headers = {"Content-Type": "application/json"}
            response.text = '{"status":true,"data":{"total":0}}'
            return response

        session = MagicMock()
        session.post.side_effect = post_side_effect
        client = LogApiClient(
            QueryConfig(token="t", output_path=Path("x.xlsx"), page_size=100),
            session=session,
        )

        pages = list(client.fetch_pages())

        self.assertEqual([], pages)
        client.close()

    def test_fetch_pages_should_treat_null_log_events_as_empty_and_stop(self) -> None:
        from unittest.mock import MagicMock

        def post_side_effect(url: str, json: dict[str, object], timeout: int) -> MagicMock:
            response = MagicMock()
            response.raise_for_status.return_value = None
            response.json.return_value = {"status": True, "data": {"logEvents": None}}
            response.headers = {"Content-Type": "application/json"}
            response.text = '{"status":true,"data":{"logEvents":null}}'
            return response

        session = MagicMock()
        session.post.side_effect = post_side_effect
        client = LogApiClient(
            QueryConfig(token="t", output_path=Path("x.xlsx"), page_size=100),
            session=session,
        )

        pages = list(client.fetch_pages())

        self.assertEqual([], pages)
        client.close()

    def test_fetch_pages_should_merge_regular_page_and_null_page(self) -> None:
        from unittest.mock import MagicMock

        pages_cursor = {"index": 0}
        normal_payload = {
            "status": True,
            "data": {
                "logEvents": [
                    {
                        "message": self._build_message("97826_dbjd", "1891463"),
                        "level": "INFO",
                    }
                ]
            },
        }
        null_payload = {"status": True, "data": {"logEvents": None}}

        def post_side_effect(url: str, json: dict[str, object], timeout: int) -> MagicMock:
            pages_cursor["index"] += 1
            payload = normal_payload if pages_cursor["index"] == 1 else null_payload
            response = MagicMock()
            response.raise_for_status.return_value = None
            response.json.return_value = payload
            response.headers = {"Content-Type": "application/json"}
            response.text = __import__("json").dumps(payload, ensure_ascii=False)
            return response

        session = MagicMock()
        session.post.side_effect = post_side_effect
        client = LogApiClient(
            QueryConfig(token="t", output_path=Path("x.xlsx"), page_size=100),
            session=session,
        )

        pages = list(client.fetch_pages())

        self.assertEqual(1, len(pages))
        self.assertEqual("97826_dbjd", pages[0][0].old_cdp_id)
        self.assertEqual("1891463", pages[0][0].bind_vip_id)
        client.close()


class HourLogExportControllerCompatibilityTests(unittest.TestCase):
    @staticmethod
    def _format_hour_label(timestamp_ms: int) -> str:
        return datetime.fromtimestamp(timestamp_ms / 1000).strftime("%Y%m%d_%H")

    def test_run_should_continue_next_hour_when_log_events_missing(self) -> None:
        first_window_start = 1_789_387_200_000
        first_window_end = 1_789_390_799_999
        second_window_start = 1_789_390_800_000
        second_window_end = 1_789_394_399_999
        first_label = self._format_hour_label(first_window_start)
        second_label = self._format_hour_label(second_window_start)
        windows = [
            HourWindow(first_window_start, first_window_end, first_label),
            HourWindow(second_window_start, second_window_end, second_label),
        ]

        class FakeWindowFactory:
            @staticmethod
            def split(start_date_ms: int, end_date_ms: int) -> list[HourWindow]:
                return list(windows)

        class FakeApplication:
            def __init__(self) -> None:
                self.calls: list[tuple[int, int, Path, int]] = []

            def __call__(self, config: QueryConfig) -> LogExportApplication:
                captured = self.calls

                class InnerApp:
                    def run(inner_self) -> Path:
                        output_path = config.output_path
                        output_path.parent.mkdir(parents=True, exist_ok=True)
                        workbook = Workbook()
                        sheet = workbook.active
                        sheet.title = "日志结果"
                        sheet.append(("Old_CdpId", "BindVipId", "source_message"))
                        if config.start_date_ms == second_window_start:
                            sheet.append(("next_dbjd", "999", "next-message"))
                        workbook.save(output_path)
                        workbook.close()
                        count_workbook = load_workbook(output_path, read_only=True)
                        try:
                            row_count = max(0, (count_workbook.active.max_row or 1) - 1)
                        finally:
                            count_workbook.close()
                        captured.append((config.start_date_ms, config.end_date_ms, output_path, row_count))
                        return output_path

                return InnerApp()

        with tempfile.TemporaryDirectory() as temp_dir:
            base_output = Path(temp_dir) / "6887-京东会员通-日志结果.xlsx"
            config = QueryConfig(
                token="custom-token",
                start_date_ms=windows[0].start_date_ms,
                end_date_ms=windows[-1].end_date_ms,
                output_path=base_output,
            )
            controller = HourLogExportController(
                base_config=config,
                window_factory=FakeWindowFactory,
                application_factory=FakeApplication(),
            )

            output_paths = controller.run()

            self.assertEqual(2, len(output_paths))
            self.assertTrue(output_paths[0].exists())
            self.assertTrue(output_paths[1].exists())
            workbook = load_workbook(output_paths[1], read_only=True)
            try:
                self.assertEqual(1, max(0, (workbook.active.max_row or 1) - 1))
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()

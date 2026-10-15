from __future__ import annotations

import argparse
import json
import logging
import re
import time
from dataclasses import dataclass, replace
from datetime import datetime
from pathlib import Path
from typing import Final

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


LOGGER = logging.getLogger("pull_logs_to_excel")

DEFAULT_TOKEN: Final = "07f58351f8c44ce8a84ff3388481f874"
DEFAULT_COOKIE: Final = "opt.authorize=267dffc44a4b4c05a4378fd3ac539eaf"
LOG_API_URL: Final[str] = "https://log.ezrpro.work/api/logs"
LOG_API_HOST: Final[str] = "log.ezrpro.work"
LOG_API_COOKIE_PATH: Final[str] = "/"
AUTH_COOKIE_NAME: Final[str] = "ezrbug_opt.authorize"
DEFAULT_APP_NAME: Final[str] = "EZP.Open.Api"
DEFAULT_CLUSTER: Final[str] = "QCloud"
DEFAULT_BRAND_ID: Final[str] = "6887"
DEFAULT_START_DATE_MS: Final[int] = 1_789_394_400_000
DEFAULT_END_DATE_MS: Final[int] = 1_789_462_799_000
DEFAULT_SEARCH_WORD: Final[str] = (
    "#loggerName:JdVipConnectionMatchVersionManager "
    "AND message:(+20240625 +_dbjd +BindVipId)"
)
DEFAULT_SORT_MODE: Final[str] = "3"
DEFAULT_PAGE_SIZE: Final[int] = 100
DEFAULT_EXCEL_BATCH_SIZE: Final[int] = DEFAULT_PAGE_SIZE
DEFAULT_TIMEOUT_SECONDS: Final[int] = 30
DEFAULT_RETRY_COUNT: Final[int] = 3
DEFAULT_RETRY_INTERVAL_SECONDS: Final[int] = 1
DEFAULT_OUTPUT_FILE_NAME: Final[str] = "6887-京东会员通-日志结果.xlsx"
EXCEL_SHEET_NAME: Final[str] = "日志结果"
EXCEL_HEADERS: Final[tuple[str, str, str]] = ("Old_CdpId", "BindVipId", "source_message")
REQUEST_HEADERS: Final[dict[str, str]] = {
    "Accept": "application/json, text/plain, */*",
    "Content-Type": "application/json;charset=UTF-8",
    "User-Agent": "pull-logs-to-excel/1.0",
}
MESSAGE_BIND_ID_PATTERN: Final[re.Pattern[str]] = re.compile(r"BindId=([^,]+)")
MESSAGE_BIND_VIP_ID_PATTERN: Final[re.Pattern[str]] = re.compile(r"BindVipId=([^,\s]+)")
HOUR_MS: Final[int] = 60 * 60 * 1000


@dataclass(frozen=True)
class HourWindow:
    start_date_ms: int
    end_date_ms: int
    hour_label: str


class HourWindowFactory:
    @staticmethod
    def _floor_to_hour(timestamp_ms: int) -> int:
        return timestamp_ms - (timestamp_ms % HOUR_MS)

    @staticmethod
    def _format_hour_label(timestamp_ms: int) -> str:
        dt_object = datetime.fromtimestamp(timestamp_ms / 1000)
        return dt_object.strftime("%Y%m%d_%H")

    @classmethod
    def split(cls, start_date_ms: int, end_date_ms: int) -> list[HourWindow]:
        if start_date_ms <= 0 or end_date_ms <= 0:
            raise ValueError("start_date_ms 和 end_date_ms 必须大于 0")
        if start_date_ms > end_date_ms:
            raise ValueError("start_date_ms 不能大于 end_date_ms")

        windows: list[HourWindow] = []
        current_start = cls._floor_to_hour(start_date_ms)
        while current_start <= end_date_ms:
            current_end = current_start + HOUR_MS - 1
            windows.append(
                HourWindow(
                    start_date_ms=max(start_date_ms, current_start),
                    end_date_ms=min(end_date_ms, current_end),
                    hour_label=cls._format_hour_label(current_start),
                )
            )
            current_start += HOUR_MS
        return windows


@dataclass(frozen=True)
class QueryConfig:
    token: str
    app_name: str = DEFAULT_APP_NAME
    cluster: str = DEFAULT_CLUSTER
    brand_id: str = DEFAULT_BRAND_ID
    start_date_ms: int = DEFAULT_START_DATE_MS
    end_date_ms: int = DEFAULT_END_DATE_MS
    search_word: str = DEFAULT_SEARCH_WORD
    sort_mode: str = DEFAULT_SORT_MODE
    page_size: int = DEFAULT_PAGE_SIZE
    excel_batch_size: int = DEFAULT_EXCEL_BATCH_SIZE
    timeout_seconds: int = DEFAULT_TIMEOUT_SECONDS
    retry_count: int = DEFAULT_RETRY_COUNT
    output_path: Path = Path(__file__).resolve().parent / DEFAULT_OUTPUT_FILE_NAME
    cookie: str = ""

    def validate(self) -> None:
        if not self.token.strip():
            raise ValueError("token 不能为空")
        if not self.brand_id.strip():
            raise ValueError("brandId 不能为空")
        if self.start_date_ms <= 0 or self.end_date_ms <= 0:
            raise ValueError("startDate 和 endDate 必须大于 0")
        if self.start_date_ms > self.end_date_ms:
            raise ValueError("startDate 不能大于 endDate")
        if self.page_size <= 0:
            raise ValueError("pageSize 必须大于 0")
        if self.excel_batch_size <= 0:
            raise ValueError("excel_batch_size 必须大于 0")
        if self.timeout_seconds <= 0:
            raise ValueError("timeout_seconds 必须大于 0")
        if self.retry_count <= 0:
            raise ValueError("retry_count 必须大于 0")


@dataclass(frozen=True)
class LogMessageRecord:
    old_cdp_id: str
    bind_vip_id: str
    source_message: str


class LogEventParser:
    def parse_message(self, source_message: str) -> LogMessageRecord:
        normalized_message = source_message.strip()
        if not normalized_message:
            raise ValueError("message 不能为空")

        bind_id_match = MESSAGE_BIND_ID_PATTERN.search(normalized_message)
        if bind_id_match is None:
            raise ValueError("message 中缺少 BindId")

        bind_id_parts = [part.strip() for part in bind_id_match.group(1).split("|")]
        if len(bind_id_parts) < 4 or not bind_id_parts[-1]:
            raise ValueError("BindId 格式不正确")

        bind_vip_id_match = MESSAGE_BIND_VIP_ID_PATTERN.search(normalized_message)
        if bind_vip_id_match is None:
            raise ValueError("message 中缺少 BindVipId")

        return LogMessageRecord(
            old_cdp_id=bind_id_parts[-1],
            bind_vip_id=bind_vip_id_match.group(1).strip(),
            source_message=normalized_message,
        )


class LogResponseParser:
    def __init__(self, event_parser: LogEventParser) -> None:
        self._event_parser = event_parser

    def parse(self, payload: dict[str, object]) -> list[LogMessageRecord]:
        data_node = payload.get("data")
        if not isinstance(data_node, dict):
            LOGGER.warning("接口响应缺少 data 节点，按空结果处理")
            return []

        log_events_node = data_node.get("logEvents")
        if log_events_node is None:
            LOGGER.warning("接口响应 logEvents 为 null，按空结果处理")
            return []
        if not isinstance(log_events_node, list):
            LOGGER.warning("接口响应 logEvents 不是列表，按空结果处理")
            return []

        result: list[LogMessageRecord] = []
        for event in log_events_node:
            if not isinstance(event, dict):
                LOGGER.warning("跳过非对象日志事件")
                continue

            message_node = event.get("message")
            if not isinstance(message_node, str) or not message_node.strip():
                LOGGER.warning("跳过缺少 message 的日志事件")
                continue

            result.append(self._event_parser.parse_message(message_node))
        return result


class LogApiClient:
    def __init__(self, config: QueryConfig, session: requests.Session | None = None) -> None:
        self._config = config
        self._session = session or requests.Session()
        self._session.headers.update(REQUEST_HEADERS)
        self._apply_cookie(config.cookie)
        self._session.cookies.set(
            AUTH_COOKIE_NAME,
            json.dumps({"token": config.token}, ensure_ascii=False, separators=(",", ":")),
            domain=LOG_API_HOST,
            path=LOG_API_COOKIE_PATH,
        )

    def fetch_pages(self):
        parser = LogResponseParser(LogEventParser())
        current_page = 0
        previous_signatures: tuple[str, ...] | None = None

        while True:
            payload = self._post_with_retry(self._build_payload(current_page))
            page_event_count = self._get_page_event_count(payload)
            rows = parser.parse(payload)
            current_signatures = tuple(row.source_message for row in rows)

            if previous_signatures is not None and current_signatures and current_signatures == previous_signatures:
                raise RuntimeError(f"分页结果重复，已停止继续拉取，page={current_page}")

            LOGGER.info("第 %s 页解析完成，新增 %s 条。", current_page, len(rows))
            if rows:
                yield self._deduplicate_current_page(rows)

            if page_event_count < self._config.page_size:
                break

            previous_signatures = current_signatures
            current_page += 1

    def close(self) -> None:
        self._session.close()

    def _build_payload(self, page: int) -> dict[str, object]:
        return {
            "appName": self._config.app_name,
            "cluster": self._config.cluster,
            "clusters": self._config.cluster,
            "hosts": "",
            "levels": "",
            "logCategories": "",
            "brandId": self._config.brand_id,
            "startDate": self._config.start_date_ms,
            "endDate": self._config.end_date_ms,
            "searchWord": self._config.search_word,
            "highlight": True,
            "histogram": True,
            "interval": 3_600_000,
            "sortMode": self._config.sort_mode,
            "page": page,
            "pageSize": self._config.page_size,
            "token": self._config.token,
        }

    def _post_with_retry(self, payload: dict[str, object]) -> dict[str, object]:
        last_error: Exception | None = None
        for current_try in range(1, self._config.retry_count + 1):
            try:
                response = self._session.post(
                    LOG_API_URL,
                    json=payload,
                    timeout=self._config.timeout_seconds,
                )
                response.raise_for_status()
                try:
                    response_json = response.json()
                except requests.JSONDecodeError as exc:
                    response_preview = response.text[:300].replace("\r", " ").replace("\n", " ")
                    raise ValueError(f"接口返回非 JSON，响应片段：{response_preview}") from exc
                if not isinstance(response_json, dict):
                    raise ValueError("接口返回不是 JSON 对象")
                if response_json.get("status") is False:
                    raise ValueError(f"接口返回失败：{response_json.get('message') or '未知错误'}")
                return response_json
            except (requests.RequestException, ValueError) as exc:
                last_error = exc
                LOGGER.warning("第 %s 次请求失败：%s", current_try, exc)
                if current_try < self._config.retry_count:
                    time.sleep(DEFAULT_RETRY_INTERVAL_SECONDS * current_try)

        raise RuntimeError(f"日志接口请求失败：{last_error}") from last_error

    @staticmethod
    def _get_page_event_count(payload: dict[str, object]) -> int:
        data_node = payload.get("data")
        if not isinstance(data_node, dict):
            LOGGER.warning("接口响应缺少 data 节点，按 0 条处理")
            return 0
        log_events = data_node.get("logEvents")
        if log_events is None:
            LOGGER.warning("接口响应 logEvents 为 null，按 0 条处理")
            return 0
        if not isinstance(log_events, list):
            LOGGER.warning("接口响应 logEvents 不是列表，按 0 条处理")
            return 0
        return len(log_events)

    def _apply_cookie(self, raw_cookie: str) -> None:
        normalized_cookie = raw_cookie.strip()
        if not normalized_cookie:
            return

        for cookie_part in normalized_cookie.split(";"):
            cookie_item = cookie_part.strip()
            if not cookie_item or "=" not in cookie_item:
                continue
            name, value = cookie_item.split("=", maxsplit=1)
            if name.strip():
                self._session.cookies.set(
                    name.strip(),
                    value.strip(),
                    domain=LOG_API_HOST,
                    path=LOG_API_COOKIE_PATH,
                )

    @staticmethod
    def _deduplicate_current_page(rows: list[LogMessageRecord]) -> list[LogMessageRecord]:
        unique_rows: list[LogMessageRecord] = []
        seen_messages: set[str] = set()
        for row in rows:
            if row.source_message in seen_messages:
                continue
            seen_messages.add(row.source_message)
            unique_rows.append(row)
        return unique_rows


class IncrementalExcelExporter:
    def __init__(self, output_path: Path) -> None:
        self._output_path = output_path

    def initialize(self) -> None:
        self._output_path.parent.mkdir(parents=True, exist_ok=True)
        if self._output_path.exists():
            self._output_path.unlink()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = EXCEL_SHEET_NAME
        sheet.append(EXCEL_HEADERS)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.column_dimensions["A"].width = 22
        sheet.column_dimensions["B"].width = 18
        sheet.column_dimensions["C"].width = 120
        sheet.freeze_panes = "A2"
        workbook.save(self._output_path)
        workbook.close()

    def append_rows(self, rows: list[LogMessageRecord]) -> None:
        if not rows:
            return

        workbook = load_workbook(self._output_path)
        try:
            sheet = workbook.active
            for row in rows:
                sheet.append((row.old_cdp_id, row.bind_vip_id, row.source_message))
            workbook.save(self._output_path)
        finally:
            workbook.close()


class LogExportApplication:
    def __init__(
        self,
        config: QueryConfig,
        client: LogApiClient | None = None,
        exporter: IncrementalExcelExporter | None = None,
    ) -> None:
        self._config = config
        self._client = client
        self._exporter = exporter

    def run(self) -> Path:
        self._config.validate()
        client = self._client or LogApiClient(self._config)
        exporter = self._exporter or IncrementalExcelExporter(self._config.output_path)
        exporter.initialize()
        pending_rows: list[LogMessageRecord] = []
        total_written = 0
        try:
            for page_rows in client.fetch_pages():
                pending_rows.extend(page_rows)
                if len(pending_rows) >= self._config.excel_batch_size:
                    exporter.append_rows(pending_rows)
                    total_written += len(pending_rows)
                    LOGGER.info("已写入一批 Excel，当前累计 %s 条。", total_written)
                    pending_rows.clear()
        finally:
            client.close()

        if pending_rows:
            exporter.append_rows(pending_rows)
            total_written += len(pending_rows)
            LOGGER.info("已写入最后一批 Excel，当前累计 %s 条。", total_written)
        LOGGER.info("导出完成，共写入 %s 条，文件：%s", total_written, self._config.output_path)
        return self._config.output_path


class HourLogExportController:
    def __init__(
        self,
        base_config: QueryConfig,
        window_factory: type[HourWindowFactory] = HourWindowFactory,
        application_factory: object = None,
    ) -> None:
        self._base_config = base_config
        self._window_factory = window_factory
        self._application_factory = application_factory or self._default_application_factory

    @staticmethod
    def _default_application_factory(config: QueryConfig) -> LogExportApplication:
        return LogExportApplication(config)

    @staticmethod
    def _build_hour_output_path(base_path: Path, hour_label: str) -> Path:
        suffix = base_path.suffix
        return base_path.with_name(f"{base_path.stem}_{hour_label}{suffix}")

    def run(self) -> list[Path]:
        self._base_config.validate()
        windows = self._window_factory.split(
            self._base_config.start_date_ms, self._base_config.end_date_ms
        )
        output_paths: list[Path] = []
        total_written = 0
        skipped_windows = 0

        for index, window in enumerate(windows, start=1):
            hour_output_path = self._build_hour_output_path(
                self._base_config.output_path, window.hour_label
            )
            try:
                LOGGER.info(
                    "开始处理第 %s/%s 小时窗口：%s，时间范围 [%s, %s]",
                    index,
                    len(windows),
                    window.hour_label,
                    window.start_date_ms,
                    window.end_date_ms,
                )
                hour_config = replace(
                    self._base_config,
                    start_date_ms=window.start_date_ms,
                    end_date_ms=window.end_date_ms,
                    output_path=hour_output_path,
                )
                application = self._application_factory(hour_config)
                application.run()

                current_rows = self._count_excel_rows(hour_output_path)
                total_written += current_rows
                LOGGER.info(
                    "小时窗口 %s 完成，写入 %s 条，累计 %s 条，文件：%s",
                    window.hour_label,
                    current_rows,
                    total_written,
                    hour_output_path,
                )
            except Exception as exc:
                skipped_windows += 1
                LOGGER.warning(
                    "小时窗口 %s 处理异常，已跳过：%s",
                    window.hour_label,
                    exc,
                )

            if hour_output_path.exists():
                output_paths.append(hour_output_path)

        LOGGER.info(
            "全部小时窗口处理完成，共生成 %s 个文件，跳过 %s 个窗口。",
            len(output_paths),
            skipped_windows,
        )
        return output_paths

    @staticmethod
    def _count_excel_rows(output_path: Path) -> int:
        workbook = load_workbook(output_path, read_only=True)
        try:
            sheet = workbook.active
            max_row = sheet.max_row or 1
            return max(0, max_row - 1)
        finally:
            workbook.close()


def configure_logging() -> None:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )


def build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="按条件拉取日志 message 并导出 Excel")
    parser.add_argument("--token", default=DEFAULT_TOKEN, help="日志接口 token")
    parser.add_argument("--cookie", default=DEFAULT_COOKIE, help="日志平台 Cookie，可选")
    parser.add_argument("--app-name", default=DEFAULT_APP_NAME, help="appName")
    parser.add_argument("--cluster", default=DEFAULT_CLUSTER, help="cluster")
    parser.add_argument("--brand-id", default=DEFAULT_BRAND_ID, help="brandId")
    parser.add_argument("--start-date-ms", type=int, default=DEFAULT_START_DATE_MS, help="开始时间毫秒时间戳")
    parser.add_argument("--end-date-ms", type=int, default=DEFAULT_END_DATE_MS, help="结束时间毫秒时间戳")
    parser.add_argument("--search-word", default=DEFAULT_SEARCH_WORD, help="检索表达式")
    parser.add_argument("--sort-mode", default=DEFAULT_SORT_MODE, help="排序模式")
    parser.add_argument("--page-size", type=int, default=DEFAULT_PAGE_SIZE, help="分页大小")
    parser.add_argument("--excel-batch-size", type=int, default=DEFAULT_EXCEL_BATCH_SIZE, help="Excel 分批写入条数")
    parser.add_argument("--timeout", type=int, default=DEFAULT_TIMEOUT_SECONDS, help="请求超时秒数")
    parser.add_argument("--retries", type=int, default=DEFAULT_RETRY_COUNT, help="重试次数")
    parser.add_argument("--output", default=DEFAULT_OUTPUT_FILE_NAME, help="输出 Excel 文件名或路径")
    return parser


def create_config_from_args(argv: list[str] | None = None) -> QueryConfig:
    args = build_argument_parser().parse_args(argv)
    output_path = Path(args.output).expanduser()
    if not output_path.is_absolute():
        output_path = Path(__file__).resolve().parent / output_path

    return QueryConfig(
        token=args.token,
        cookie=args.cookie,
        app_name=args.app_name,
        cluster=args.cluster,
        brand_id=args.brand_id,
        start_date_ms=args.start_date_ms,
        end_date_ms=args.end_date_ms,
        search_word=args.search_word,
        sort_mode=args.sort_mode,
        page_size=args.page_size,
        excel_batch_size=args.excel_batch_size,
        timeout_seconds=args.timeout,
        retry_count=args.retries,
        output_path=output_path,
    )


def main() -> int:
    configure_logging()
    try:
        HourLogExportController(create_config_from_args()).run()
        return 0
    except Exception as exc:
        LOGGER.exception("执行失败：%s", exc)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())

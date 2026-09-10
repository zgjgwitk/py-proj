r"""
功能说明
========

1. 通过 busi-log 日志系统查询日志，查询方式参考：
   D:\Github\py-proj\京东会员通bugfix\20260615-unionid\jd_unionid_export.py
2. 根据查询结果解析日志字段，并生成 Excel 文件，默认输出到：
   D:\Github\py-proj\京东灰度接口\file
3. Excel 必须包含原始字段：`httpPath`、`requestContent`。
4. 根据 `httpPath` 的值解析 `requestContent` JSON，并将字段写入对应 Excel 列。
   4.1 `/api/member/AdjustPoint`
       - 解析 requestContent 的一级字段。
       - `extend` 继续解析为 `extend.businessId`、`extend.omid` 等独立列。
   4.2 `/api/member/UpdateMemberData`
       - 解析 requestContent 的一级字段。
       - `memberData`、`memberData.point` 写入独立列。
       - `memberData.extend` 继续解析为 `memberData.extend.jdPoint` 等独立列。
   4.3 `/api/member/MemberRegister`
       - 解析 requestContent 的一级字段。
       - `extend` 继续解析为 `extend.point` 等独立列。
5. 严格校验规则：
   - 发现未列出的 httpPath 时，终止程序并抛出明确错误。
   - 发现示例中未定义的字段时，终止程序并抛出明确错误。
   - 缺少 httpPath、requestContent，或 requestContent / 嵌套 JSON 格式错误时，终止程序。
6. 后续需求变更时，请优先同步修改本说明和下方 ROUTE_SCHEMAS 字段白名单，
   再调整对应解析逻辑。
=== 追加逻辑1
7. 生成excel完成后, 需要对excel中的数据进行过滤去重, 规则是:
   排除掉[日志时间,sessionId,traceId,requestContent,timestamp,token,occurTime]
以外的字段都相同的, 只要保留一个即可
8. 支持仅输入 Excel 文件路径执行第 7 步：
   使用 `--deduplicate-input` 指定源文件时，不查询日志、不修改源文件，
   默认在源文件同目录生成追加 `_去重` 的新 Excel 文件。
"""

from __future__ import annotations

import argparse
import json
import logging
import time
from dataclasses import dataclass
from datetime import datetime, time as datetime_time, timedelta
from pathlib import Path
from typing import Final
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


LOG_API_URL: Final = "https://log.ezrpro.work/api/logs"
DEFAULT_COOKIE: Final = (
    "opt.authorize=267dffc44a4b4c05a4378fd3ac539eaf; "
    "jwt_token=eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9."
    "eyJ1c2VyaWQiOjI3OSwibmFtZSI6InNzb2xvZ2luIiwidXNlcm5hbWUiOiJaaGFuZ0ppbmdXZWkiLCJleHBpciI6MTc4ODg1NDg1Nn0."
    "YStSFYMUAUyadNepZXn3wzIQ2A43qL4-GcurdDT5PeY"
)
DEFAULT_START_TIME: Final = "2026-09-08 14:29:00"
DEFAULT_END_TIME: Final = "2026-09-08 15:51:05"
# brandId 是可选过滤条件；为空时查询全部品牌，仍可通过 --brand-id 指定。
DEFAULT_BRAND_ID: Final = ""
DEFAULT_SEARCH_WORD: Final = (
    "#requestContent:(+omid +100000000037737) "
    "AND message:(/api/member/) AND duration:>1"
)
DEFAULT_TOKEN: Final = "07f58351f8c44ce8a84ff3388481f874"
DEFAULT_APP_NAME: Final = "EZR.OP.ThirdPartyProxy.ApiHost"
DEFAULT_CLUSTER: Final = "QCloud"
DEFAULT_LOG_CATEGORIES: Final = "HttpServer"
DEFAULT_TIMEZONE: Final = "Asia/Shanghai"
DEFAULT_PAGE_SIZE: Final = 100
DEFAULT_TIMEOUT_SECONDS: Final = 30
DEFAULT_RETRY_COUNT: Final = 3
RETRY_INTERVAL_SECONDS: Final = 1
HISTOGRAM_INTERVAL_MS: Final = 3_600_000
MILLISECONDS_PER_SECOND: Final = 1_000
SORT_MODE: Final = "3"
EXCEL_CELL_LENGTH_LIMIT: Final = 32_767
EXCEL_HEADER_FILL: Final = "D9EAF7"
MIN_COLUMN_WIDTH: Final = 12
MAX_COLUMN_WIDTH: Final = 60
OUTPUT_DIRECTORY_NAME: Final = "file"
OUTPUT_SHEET_NAME: Final = "请求数据"
DATETIME_FORMAT: Final = "%Y-%m-%d %H:%M:%S"

# None 表示普通字段；字典表示需要继续校验并展开的嵌套 JSON 对象。
# 字段白名单来自需求中的三个示例。出现白名单外字段时必须终止程序，
# 避免脚本在字段含义未知的情况下生成不完整数据。
ROUTE_SCHEMAS: Final[dict[str, dict[str, object]]] = {
    "/api/member/AdjustPoint": {
        "account": None,
        "appkey": None,
        "brandId": None,
        "changeType": None,
        "content": None,
        "extend": {
            "businessId": None,
            "omid": None,
        },
        "occurTime": None,
        "platform": None,
        "point": None,
        "pointType": None,
        "timestamp": None,
        "token": None,
    },
    "/api/member/UpdateMemberData": {
        "MixPhone": None,
        "account": None,
        "appkey": None,
        "bind_status": None,
        "brandId": None,
        "cardStatus": None,
        "jd_bind_time": None,
        "memberData": {
            "extend": {
                "jdPoint": None,
            },
            "point": None,
        },
        "omid": None,
        "platform": None,
        "timestamp": None,
        "token": None,
        "totalOrderCount": None,
        "totalOrderPrice": None,
    },
    "/api/member/MemberRegister": {
        "MixPhone": None,
        "appkey": None,
        "brandId": None,
        "extend": {
            "point": None,
        },
        "jd_bind_time": None,
        "omid": None,
        "platform": None,
        "timestamp": None,
        "token": None,
    },
}

LOG_METADATA_HEADERS: Final = (
    "日志时间",
    "sessionId",
    "traceId",
    "statusCode",
    "httpPath",
    "requestContent",
)
DEDUPLICATION_EXCLUDED_HEADERS: Final = frozenset(
    {
        "日志时间",
        "sessionId",
        "traceId",
        "requestContent",
        "timestamp",
        "token",
        "occurTime",
    }
)

LOGGER = logging.getLogger("jd_gray_log_exporter")


@dataclass(frozen=True)
class QueryConfig:
    """日志查询和文件导出配置。"""

    start_at: datetime
    end_at: datetime
    timezone: ZoneInfo
    output_path: Path
    app_name: str
    cluster: str
    log_categories: str
    brand_id: str
    search_word: str
    token: str
    cookie: str
    page_size: int = DEFAULT_PAGE_SIZE
    timeout_seconds: int = DEFAULT_TIMEOUT_SECONDS
    retry_count: int = DEFAULT_RETRY_COUNT

    def validate(self) -> None:
        """在调用外部接口前校验配置，避免产生无效请求。"""

        if self.start_at > self.end_at:
            raise ValueError("开始时间不能晚于结束时间")
        if not self.token.strip():
            raise ValueError("token 不能为空")
        if not self.cookie.strip():
            raise ValueError("cookie 不能为空")
        if self.page_size <= 0:
            raise ValueError("page-size 必须大于 0")
        if self.timeout_seconds <= 0:
            raise ValueError("timeout 必须大于 0")
        if self.retry_count <= 0:
            raise ValueError("retries 必须大于 0")


@dataclass(frozen=True)
class QueryWindow:
    """单个自然日内的查询时间窗口。"""

    start_at: datetime
    end_at: datetime

    @property
    def start_ms(self) -> int:
        return int(self.start_at.timestamp() * MILLISECONDS_PER_SECOND)

    @property
    def end_ms(self) -> int:
        return int(self.end_at.timestamp() * MILLISECONDS_PER_SECOND)


class QueryWindowFactory:
    """把跨日时间范围拆分成自然日窗口。"""

    @staticmethod
    def create(start_at: datetime, end_at: datetime) -> list[QueryWindow]:
        windows: list[QueryWindow] = []
        current_start = start_at

        while current_start <= end_at:
            day_end = datetime.combine(
                current_start.date(),
                datetime_time(23, 59, 59, 999000),
                current_start.tzinfo,
            )
            current_end = min(day_end, end_at)
            windows.append(QueryWindow(current_start, current_end))
            current_start = current_end + timedelta(milliseconds=1)

        return windows


class PayloadParser:
    """按接口路径白名单校验并展开 requestContent。"""

    def parse(self, event: object, page: int, row_number: int) -> dict[str, object]:
        location = self._build_location(event, page, row_number)
        if not isinstance(event, dict):
            raise ValueError(f"日志事件不是 JSON 对象：{location}")

        http_path = self._required_string(event, "httpPath", location)
        schema = ROUTE_SCHEMAS.get(http_path)
        if schema is None:
            raise ValueError(f"发现未支持的 httpPath={http_path!r}：{location}")

        request_content = self._required_string(event, "requestContent", location)
        if len(request_content) > EXCEL_CELL_LENGTH_LIMIT:
            raise ValueError(
                f"requestContent 超过 Excel 单元格长度限制：{location}"
            )

        payload = self._parse_json_object(
            request_content,
            "requestContent",
            location,
        )
        parsed_fields: dict[str, object] = {}
        self._validate_and_flatten(
            payload,
            schema,
            prefix="",
            output=parsed_fields,
            location=location,
        )

        return {
            "日志时间": self._parse_timestamp(event.get("createdOn")),
            "sessionId": str(event.get("sessionId") or ""),
            "traceId": str(event.get("traceId") or ""),
            "statusCode": str(event.get("statusCode") or ""),
            "httpPath": http_path,
            "requestContent": request_content,
            **parsed_fields,
        }

    def _validate_and_flatten(
        self,
        data: dict[str, object],
        schema: dict[str, object],
        prefix: str,
        output: dict[str, object],
        location: str,
    ) -> None:
        """递归校验字段白名单，并使用点号生成嵌套字段列名。"""

        unknown_fields = sorted(set(data) - set(schema))
        if unknown_fields:
            object_name = prefix.removesuffix(".") or "requestContent"
            raise ValueError(
                f"{object_name} 出现未定义字段 {unknown_fields}：{location}"
            )

        for field_name, value in data.items():
            column_name = f"{prefix}{field_name}"
            nested_schema = schema[field_name]

            if nested_schema is None:
                if isinstance(value, (dict, list)):
                    raise ValueError(
                        f"字段 {column_name} 应为普通值，实际为嵌套结构："
                        f"{location}"
                    )
                output[column_name] = value
                continue

            if not isinstance(nested_schema, dict):
                raise TypeError(f"字段 {column_name} 的内部解析规则配置错误")

            nested_data = self._to_nested_object(value, column_name, location)
            # 同时保留嵌套对象本身，便于和原始日志核对。
            output[column_name] = json.dumps(
                nested_data,
                ensure_ascii=False,
                separators=(",", ":"),
            )
            self._validate_and_flatten(
                nested_data,
                nested_schema,
                prefix=f"{column_name}.",
                output=output,
                location=location,
            )

    def _to_nested_object(
        self,
        value: object,
        field_name: str,
        location: str,
    ) -> dict[str, object]:
        """兼容嵌套字段为对象或 JSON 字符串的两种日志格式。"""

        if isinstance(value, dict):
            return value
        return self._parse_json_object(value, field_name, location)

    @staticmethod
    def _parse_json_object(
        value: object,
        field_name: str,
        location: str,
    ) -> dict[str, object]:
        if not isinstance(value, str) or not value.strip():
            raise ValueError(f"字段 {field_name} 必须是 JSON 对象：{location}")

        try:
            parsed_value = json.loads(value)
        except json.JSONDecodeError as exc:
            raise ValueError(
                f"字段 {field_name} JSON 无效：{exc.msg}（位置 {exc.pos}）："
                f"{location}"
            ) from exc

        if not isinstance(parsed_value, dict):
            raise ValueError(
                f"字段 {field_name} 的 JSON 根节点必须是对象：{location}"
            )
        return parsed_value

    @staticmethod
    def _required_string(
        event: dict[str, object],
        field_name: str,
        location: str,
    ) -> str:
        value = event.get(field_name)
        if not isinstance(value, str) or not value.strip():
            raise ValueError(f"日志缺少必填字段 {field_name}：{location}")
        return value.strip()

    @staticmethod
    def _parse_timestamp(value: object) -> int:
        if value is None or isinstance(value, bool):
            return 0
        try:
            return int(value)
        except (TypeError, ValueError):
            return 0

    @staticmethod
    def _build_location(event: object, page: int, row_number: int) -> str:
        session_id = event.get("sessionId") if isinstance(event, dict) else ""
        return (
            f"page={page}, row={row_number}, "
            f"sessionId={session_id or ''}"
        )


class LogApiClient:
    """封装 busi-log 查询、分页保护和请求重试。"""

    def __init__(self, config: QueryConfig, parser: PayloadParser) -> None:
        self._config = config
        self._parser = parser
        self._session = requests.Session()
        self._session.headers.update(
            {
                "Accept": "application/json, text/plain, */*",
                "Content-Type": "application/json;charset=UTF-8",
                "Cookie": config.cookie,
                "User-Agent": "JdGrayLogExporter/1.0",
            }
        )

    def close(self) -> None:
        self._session.close()

    def fetch(self, window: QueryWindow) -> list[dict[str, object]]:
        page = 0
        received_count = 0
        api_total: int | None = None
        previous_page_signature: tuple[str, ...] | None = None
        records: list[dict[str, object]] = []

        while True:
            response_data = self._post_json(self._build_payload(window, page))
            data_node = response_data.get("data")
            if not isinstance(data_node, dict):
                raise ValueError("日志接口响应缺少 data 对象")

            if api_total is None:
                api_total = self._to_optional_int(data_node.get("total"))

            events = data_node.get("logEvents") or []
            if not isinstance(events, list):
                raise ValueError("日志接口响应 data.logEvents 不是数组")

            LOGGER.info(
                "第 %s 页返回 %s 条，接口总数 %s",
                page,
                len(events),
                api_total if api_total is not None else "未知",
            )
            if not events:
                break

            # 防止异常接口忽略 page 参数而持续返回相同的数据。
            page_signature = tuple(repr(event) for event in events)
            if page_signature == previous_page_signature:
                raise RuntimeError(
                    f"接口连续返回相同分页，已终止查询：page={page}"
                )
            previous_page_signature = page_signature

            for row_number, event in enumerate(events, start=1):
                records.append(self._parser.parse(event, page, row_number))
            received_count += len(events)

            if api_total is not None and received_count >= api_total:
                break
            if api_total is None and len(events) < self._config.page_size:
                break
            page += 1

        return records

    def _build_payload(
        self,
        window: QueryWindow,
        page: int,
    ) -> dict[str, object]:
        payload = {
            "appName": self._config.app_name,
            "cluster": self._config.cluster,
            "clusters": self._config.cluster,
            "hosts": "",
            "levels": "",
            "logCategories": self._config.log_categories,
            "brandId": self._config.brand_id,
            "startDate": window.start_ms,
            "endDate": window.end_ms,
            "searchWord": self._config.search_word,
            "highlight": True,
            "histogram": True,
            "interval": HISTOGRAM_INTERVAL_MS,
            "sortMode": SORT_MODE,
            "page": page,
            "pageSize": self._config.page_size,
            "token": self._config.token,
        }
        # 只记录可公开核对的查询条件，不输出 Cookie 和 token 等敏感信息。
        LOGGER.info(
            "组装日志查询参数：appName=%s，brandId=%r，logCategories=%r，"
            "startDate=%s，endDate=%s，page=%s，pageSize=%s",
            payload["appName"],
            payload["brandId"],
            payload["logCategories"],
            payload["startDate"],
            payload["endDate"],
            payload["page"],
            payload["pageSize"],
        )
        return payload

    def _post_json(self, payload: dict[str, object]) -> dict[str, object]:
        last_error: Exception | None = None

        for attempt in range(1, self._config.retry_count + 1):
            try:
                response = self._session.post(
                    LOG_API_URL,
                    json=payload,
                    timeout=self._config.timeout_seconds,
                )
                response.raise_for_status()
                response_data = response.json()

                if not isinstance(response_data, dict):
                    raise ValueError("日志接口响应不是 JSON 对象")
                if response_data.get("status") is False:
                    message = response_data.get("message") or "未知错误"
                    raise RuntimeError(f"日志接口返回失败：{message}")
                return response_data
            except (
                requests.RequestException,
                requests.JSONDecodeError,
                ValueError,
                RuntimeError,
            ) as exc:
                last_error = exc
                LOGGER.warning(
                    "日志接口请求失败，第 %s/%s 次：%s",
                    attempt,
                    self._config.retry_count,
                    exc,
                )
                if attempt < self._config.retry_count:
                    time.sleep(RETRY_INTERVAL_SECONDS * attempt)

        raise RuntimeError(f"日志接口请求最终失败：{last_error}") from last_error

    @staticmethod
    def _to_optional_int(value: object) -> int | None:
        if value is None or isinstance(value, bool):
            return None
        try:
            return int(value)
        except (TypeError, ValueError):
            return None


class ExcelExporter:
    """生成统一字段表头，并通过临时文件安全导出 Excel。"""

    @staticmethod
    def build_headers() -> tuple[str, ...]:
        headers = list(LOG_METADATA_HEADERS)
        for schema in ROUTE_SCHEMAS.values():
            ExcelExporter._append_schema_headers(schema, "", headers)
        # 不同接口存在同名字段，按首次出现顺序去重。
        return tuple(dict.fromkeys(headers))

    @staticmethod
    def _append_schema_headers(
        schema: dict[str, object],
        prefix: str,
        headers: list[str],
    ) -> None:
        for field_name, nested_schema in schema.items():
            column_name = f"{prefix}{field_name}"
            headers.append(column_name)
            if isinstance(nested_schema, dict):
                ExcelExporter._append_schema_headers(
                    nested_schema,
                    f"{column_name}.",
                    headers,
                )

    def export(
        self,
        output_path: Path,
        records: list[dict[str, object]],
        timezone: ZoneInfo,
    ) -> None:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = OUTPUT_SHEET_NAME
        headers = self.build_headers()
        sheet.append(headers)

        for record in records:
            row = [
                self._to_excel_value(record.get(header), header, timezone)
                for header in headers
            ]
            sheet.append(row)

        self._deduplicate_sheet(sheet, headers)
        self._format_sheet(sheet, headers)

        # 先写临时文件，校验或保存失败时不会覆盖已有结果。
        temporary_path = output_path.with_name(
            f"{output_path.stem}.tmp{output_path.suffix}"
        )
        try:
            workbook.save(temporary_path)
            try:
                temporary_path.replace(output_path)
            except PermissionError as exc:
                raise PermissionError(
                    f"无法覆盖 Excel 文件 {output_path}，请先关闭正在打开的文件，"
                    "并确认当前用户具有写入权限"
                ) from exc
        finally:
            workbook.close()
            if temporary_path.exists():
                temporary_path.unlink()

    @staticmethod
    def _deduplicate_sheet(sheet: object, headers: tuple[str, ...]) -> None:
        """按需求排除指定字段后去重，保留首次出现的记录。"""

        comparison_indexes = [
            index
            for index, header in enumerate(headers)
            if header not in DEDUPLICATION_EXCLUDED_HEADERS
        ]
        seen_keys: set[tuple[object, ...]] = set()
        duplicate_row_numbers: list[int] = []

        for row_number, row in enumerate(
            sheet.iter_rows(min_row=2),  # type: ignore[attr-defined]
            start=2,
        ):
            comparison_key = tuple(row[index].value for index in comparison_indexes)
            if comparison_key in seen_keys:
                duplicate_row_numbers.append(row_number)
                continue
            seen_keys.add(comparison_key)

        # 从后向前删除，避免删除前面的行后导致行号偏移。
        for row_number in reversed(duplicate_row_numbers):
            sheet.delete_rows(row_number, 1)  # type: ignore[attr-defined]

        if duplicate_row_numbers:
            LOGGER.info(
                "Excel 去重完成：删除 %s 条重复记录，保留 %s 条",
                len(duplicate_row_numbers),
                len(seen_keys),
            )

    @staticmethod
    def _to_excel_value(
        value: object,
        header: str,
        timezone: ZoneInfo,
    ) -> object:
        if header == "日志时间" and isinstance(value, int) and value > 0:
            try:
                created_at = datetime.fromtimestamp(
                    value / MILLISECONDS_PER_SECOND,
                    timezone,
                )
                return created_at.strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
            except (OverflowError, OSError, ValueError):
                LOGGER.warning("createdOn 时间戳无效：%s", value)
                return str(value)

        if value is None:
            return ""
        if isinstance(value, (dict, list)):
            return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
        return value

    @staticmethod
    def _format_sheet(sheet: object, headers: tuple[str, ...]) -> None:
        # openpyxl 的 Worksheet 缺少稳定的公开类型入口，此处接收实际工作表对象。
        for cell in sheet[1]:  # type: ignore[index]
            cell.font = Font(bold=True)
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor=EXCEL_HEADER_FILL,
            )
            cell.alignment = Alignment(horizontal="center", vertical="center")

        sheet.freeze_panes = "A2"  # type: ignore[attr-defined]
        sheet.auto_filter.ref = sheet.dimensions  # type: ignore[attr-defined]
        for column_index, header in enumerate(headers, start=1):
            column_width = min(
                max(len(header) + 2, MIN_COLUMN_WIDTH),
                MAX_COLUMN_WIDTH,
            )
            sheet.column_dimensions[  # type: ignore[attr-defined]
                get_column_letter(column_index)
            ].width = column_width


class ExcelFileDeduplicator:
    """读取已有 Excel 并将去重结果另存为新文件。"""

    def deduplicate(
        self,
        input_path: Path,
        output_path: Path | None = None,
    ) -> Path:
        if not input_path.exists():
            raise FileNotFoundError(f"源 Excel 文件不存在：{input_path}")
        if input_path.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise ValueError("源文件必须是 .xlsx 或 .xlsm 格式")

        source_path = input_path.resolve()
        target_path = output_path or self._build_output_path(source_path)
        target_path = target_path.resolve()
        if source_path == target_path:
            raise ValueError("去重输出文件不能与源 Excel 文件相同")
        target_path.parent.mkdir(parents=True, exist_ok=True)

        # 以可编辑模式加载，保留源文件的工作表、样式和其他内容；只写入新文件。
        workbook = load_workbook(source_path, keep_vba=source_path.suffix.lower() == ".xlsm")
        try:
            sheet = workbook.active
            headers = tuple(
                str(cell.value).strip() if cell.value is not None else ""
                for cell in sheet[1]
            )
            if not any(headers):
                raise ValueError("源 Excel 第一行没有有效表头")

            before_count = max(sheet.max_row - 1, 0)
            ExcelExporter._deduplicate_sheet(sheet, headers)
            after_count = max(sheet.max_row - 1, 0)
            self._save_as_new_file(workbook, target_path)
            LOGGER.info(
                "源 Excel 去重完成：原始 %s 条，保留 %s 条，输出：%s",
                before_count,
                after_count,
                target_path,
            )
            return target_path
        finally:
            workbook.close()

    @staticmethod
    def _build_output_path(source_path: Path) -> Path:
        return source_path.with_name(f"{source_path.stem}_去重{source_path.suffix}")

    @staticmethod
    def _save_as_new_file(workbook: Workbook, output_path: Path) -> None:
        temporary_path = output_path.with_name(
            f"{output_path.stem}.tmp{output_path.suffix}"
        )
        try:
            workbook.save(temporary_path)
            temporary_path.replace(output_path)
        except PermissionError as exc:
            raise PermissionError(
                f"无法写入新的 Excel 文件 {output_path}，请关闭同名文件并检查写入权限"
            ) from exc
        finally:
            if temporary_path.exists():
                temporary_path.unlink()


class GrayLogExportApplication:
    """协调日志查询、数据去重和 Excel 导出。"""

    def __init__(self, config: QueryConfig) -> None:
        self._config = config

    def run(self) -> int:
        self._config.validate()
        windows = QueryWindowFactory.create(
            self._config.start_at,
            self._config.end_at,
        )
        records: list[dict[str, object]] = []
        seen_keys: set[tuple[object, object, object, object]] = set()
        client = LogApiClient(self._config, PayloadParser())

        try:
            for window in windows:
                LOGGER.info(
                    "开始查询时间窗口：%s 至 %s",
                    window.start_at,
                    window.end_at,
                )
                for record in client.fetch(window):
                    unique_key = self._build_unique_key(record)
                    if unique_key in seen_keys:
                        LOGGER.warning(
                            "跳过重复日志：sessionId=%s",
                            record.get("sessionId") or "",
                        )
                        continue
                    seen_keys.add(unique_key)
                    records.append(record)
        finally:
            client.close()

        # 所有数据完成严格校验后才创建结果文件。
        ExcelExporter().export(
            self._config.output_path,
            records,
            self._config.timezone,
        )
        LOGGER.info(
            "导出完成：共 %s 条，文件：%s",
            len(records),
            self._config.output_path,
        )
        return 0

    @staticmethod
    def _build_unique_key(
        record: dict[str, object],
    ) -> tuple[object, object, object, object]:
        return (
            record.get("日志时间"),
            record.get("sessionId"),
            record.get("httpPath"),
            record.get("requestContent"),
        )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="查询京东灰度接口日志并导出 Excel"
    )
    parser.add_argument("--start-time", default=DEFAULT_START_TIME)
    parser.add_argument("--end-time", default=DEFAULT_END_TIME)
    parser.add_argument(
        "--brand-id",
        default=DEFAULT_BRAND_ID,
        help="可选的日志品牌过滤条件，留空时查询全部品牌",
    )
    parser.add_argument("--search-word", default=DEFAULT_SEARCH_WORD)
    parser.add_argument("--token", default=DEFAULT_TOKEN)
    parser.add_argument("--cookie", default=DEFAULT_COOKIE)
    parser.add_argument("--app-name", default=DEFAULT_APP_NAME)
    parser.add_argument("--cluster", default=DEFAULT_CLUSTER)
    parser.add_argument("--log-categories", default=DEFAULT_LOG_CATEGORIES)
    parser.add_argument("--timezone", default=DEFAULT_TIMEZONE)
    parser.add_argument("--page-size", type=int, default=DEFAULT_PAGE_SIZE)
    parser.add_argument("--timeout", type=int, default=DEFAULT_TIMEOUT_SECONDS)
    parser.add_argument("--retries", type=int, default=DEFAULT_RETRY_COUNT)
    parser.add_argument("--output", default="", help="Excel 输出路径")
    parser.add_argument(
        "--deduplicate-input",
        default="",
        help="仅对指定源 Excel 执行去重，不查询日志且不修改源文件",
    )
    return parser.parse_args()


def create_config(args: argparse.Namespace) -> QueryConfig:
    try:
        timezone = ZoneInfo(args.timezone)
    except ZoneInfoNotFoundError as exc:
        raise ValueError(f"无效时区：{args.timezone}") from exc

    try:
        start_at = datetime.strptime(args.start_time, DATETIME_FORMAT).replace(
            tzinfo=timezone
        )
        end_at = datetime.strptime(args.end_time, DATETIME_FORMAT).replace(
            tzinfo=timezone
        )
    except ValueError as exc:
        raise ValueError(f"时间格式必须为 {DATETIME_FORMAT}") from exc

    # 与浏览器请求保持一致：结束时间使用用户输入秒数对应的毫秒值。
    output_path = build_output_path(args.output, start_at, end_at)

    return QueryConfig(
        start_at=start_at,
        end_at=end_at,
        timezone=timezone,
        output_path=output_path,
        app_name=args.app_name,
        cluster=args.cluster,
        log_categories=args.log_categories,
        brand_id=args.brand_id,
        search_word=args.search_word,
        token=args.token,
        cookie=args.cookie,
        page_size=args.page_size,
        timeout_seconds=args.timeout,
        retry_count=args.retries,
    )


def build_output_path(
    configured_path: str,
    start_at: datetime,
    end_at: datetime,
) -> Path:
    if configured_path.strip():
        return Path(configured_path).expanduser().resolve()

    file_name = (
        f"京东灰度接口日志_{start_at:%Y%m%d_%H%M%S}_"
        f"{end_at:%Y%m%d_%H%M%S}.xlsx"
    )
    return Path(__file__).resolve().parent / OUTPUT_DIRECTORY_NAME / file_name


def configure_logging() -> None:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )


def main() -> int:
    configure_logging()
    try:
        args = parse_args()
        if args.deduplicate_input.strip():
            input_path = Path(args.deduplicate_input).expanduser().resolve()
            output_path = (
                Path(args.output).expanduser().resolve()
                if args.output.strip()
                else None
            )
            result_path = ExcelFileDeduplicator().deduplicate(
                input_path,
                output_path,
            )
            LOGGER.info("独立去重完成：%s", result_path)
            return 0

        config = create_config(args)
        return GrayLogExportApplication(config).run()
    except KeyboardInterrupt:
        LOGGER.warning("用户中断执行")
        return 130
    except Exception as exc:
        LOGGER.exception("执行失败：%s", exc)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())

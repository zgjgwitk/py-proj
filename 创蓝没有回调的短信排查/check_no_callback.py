"""
排查没有回调的品牌_渠道
从 ES 索引中查询 resStatus 仅有 [1, 2] 的数据,按 brandId, serviceType, resStatus 分组并导出 Excel
"""

from elasticsearch import Elasticsearch
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import logging
import sys
from pathlib import Path

# ES 地址
ES_HOST = "http://192.168.12.124:88/@qcloud:base.es.biz-172.21.65.197:9200/"

# 待查询的索引集合
INDEX_NAMES = [
    # "esmsgsms2601",
    # "esmsgsms2602",
    # "esmsgsms2603",
    # "esmsgsms2604",
    # "esmsgsms2605",
    # "esmsgsms2606",
    "esmsgsms2607",
]

# 需要校验的 resStatus 取值范围
TARGET_RES_STATUS = [1, 2]

# 输出目录
OUTPUT_DIR = Path(__file__).parent

# 日志配置
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[logging.StreamHandler(sys.stdout)],
)
logger = logging.getLogger(__name__)


def query_es_group_by(es: Elasticsearch, index_name: str, target_statuses: list[int]) -> dict:
    """
    在指定索引上聚合,统计每个 (brandId, serviceType, resStatus) 分组的文档数。
    过滤条件: 该 brandId + serviceType 分组下出现的 resStatus 集合仅为 target_statuses。
    """
    # 1) 先按 brandId, serviceType, resStatus 聚合
    agg_body = {
        "size": 0,
        "query": {
            "bool": {
                "must_not": [
                    {"term": {"smsChan": 100}}
                ]
            }
        },
        "aggs": {
            "by_brand_service": {
                "terms": {"field": "brandId", "size": 10000},
                "aggs": {
                    "by_service": {
                        "terms": {"field": "serviceType", "size": 1000},
                        "aggs": {
                            "by_status": {
                                "terms": {"field": "resStatus", "size": 20},
                            }
                        },
                    }
                },
            }
        },
    }

    logger.info("开始查询索引: %s", index_name)
    response = es.search(index=index_name, body=agg_body, ignore_unavailable=True)

    # 2) 遍历聚合,仅保留 resStatus 集合恰为 target_statuses 的分组
    result: dict[tuple, dict] = {}
    aggs = response.get("aggregations", {}).get("by_brand_service", {}).get("buckets", [])

    for brand_bucket in aggs:
        brand_id = brand_bucket.get("key")
        for service_bucket in brand_bucket.get("by_service", {}).get("buckets", []):
            service_type = service_bucket.get("key")
            status_buckets = service_bucket.get("by_status", {}).get("buckets", [])
            status_keys = {int(b.get("key")) for b in status_buckets}

            # 仅当出现的 resStatus 完全等于目标集合时保留
            if status_keys != set(target_statuses):
                continue

            # 把每个目标状态对应的文档数收集起来
            status_counts: dict[int, int] = {}
            for b in status_buckets:
                status_counts[int(b.get("key"))] = b.get("doc_count", 0)

            for status, count in status_counts.items():
                key = (brand_id, service_type, status)
                result[key] = {
                    "brandId": brand_id,
                    "serviceType": service_type,
                    "resStatus": status,
                    "docCount": count,
                }

    logger.info("索引 %s 匹配分组数: %d", index_name, len(result))
    return result


def export_to_excel_per_index(data_by_index: dict[str, dict], output_dir: Path) -> None:
    """
    将结果按索引拆分为多个 Excel 文件,每个索引生成一个独立的 xlsx 文件。
    """
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")

    columns = ["brandId", "serviceType", "resStatus", "文档数"]

    for index_name, group_data in data_by_index.items():
        # 确保输出目录存在
        output_dir.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = index_name[:31]  # Excel sheet 名最长 31 字符

        # 写表头
        for col_idx, col_name in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align

        # 写数据
        row = 2
        sorted_items = sorted(
            group_data.values(),
            key=lambda item: (str(item["brandId"]), str(item["serviceType"]), item["resStatus"]),
        )
        for item in sorted_items:
            ws.cell(row=row, column=1, value=item["brandId"])
            ws.cell(row=row, column=2, value=item["serviceType"])
            ws.cell(row=row, column=3, value=item["resStatus"])
            ws.cell(row=row, column=4, value=item["docCount"])
            row += 1

        # 调整列宽
        for col_idx in range(1, len(columns) + 1):
            ws.column_dimensions[chr(64 + col_idx)].width = 22

        output_path = output_dir / f"排查没有回调的品牌_渠道_{index_name}.xlsx"
        wb.save(output_path)
        logger.info("Excel 已导出: %s (记录数: %d)", output_path, len(sorted_items))


def main() -> None:
    es = Elasticsearch(ES_HOST, request_timeout=60)
    try:
        # 健康检查
        if not es.ping():
            logger.error("ES 连接失败: %s", ES_HOST)
            return
        logger.info("ES 连接成功: %s", ES_HOST)
    except Exception as exc:
        logger.error("ES 连接异常: %s", exc)
        return

    data_by_index: dict[str, dict] = {}
    for index_name in INDEX_NAMES:
        try:
            group_data = query_es_group_by(es, index_name, TARGET_RES_STATUS)
            data_by_index[index_name] = group_data
        except Exception as exc:
            logger.error("查询索引 %s 失败: %s", index_name, exc)
            data_by_index[index_name] = {}

    output_path = OUTPUT_DIR / "排查没有回调的品牌_渠道"
    export_to_excel_per_index(data_by_index, output_path)
    logger.info("处理完成,共处理索引数: %d", len(data_by_index))


if __name__ == "__main__":
    main()

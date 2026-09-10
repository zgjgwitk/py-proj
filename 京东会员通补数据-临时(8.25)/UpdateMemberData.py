from __future__ import annotations

import argparse
import json
import logging
import time
from dataclasses import dataclass
from datetime import date, datetime, time as datetime_time, timedelta
from pathlib import Path
from typing import Final
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


LOG_API_URL: Final = "https://log.ezrpro.work/api/logs"
LOG_API_COOKIE_DOMAIN: Final = "log.ezrpro.work"
LOG_API_COOKIE_PATH: Final = "/"
LOG_API_AUTH_COOKIE_NAME: Final = "ezrbug_opt.authorize"
DEFAULT_BRAND_IDS: Final = (
    "190", "7097", "7003", "7091", "7126",
    "7095", "5417", "1752", "7054", "7135"
)
DEFAULT_START_DATE: Final = "2026-08-25 19:30:00"
DEFAULT_END_DATE: Final = "2026-08-25 22:10:00"
DEFAULT_TOKEN: Final = "07f58351f8c44ce8a84ff3388481f874"
DEFAULT_COOKIE: Final = "opt.authorize=267dffc44a4b4c05a4378fd3ac539eaf"
DEFAULT_TIMEZONE: Final = "Asia/Shanghai"
DEFAULT_PAGE_SIZE: Final = 100
DEFAULT_TIMEOUT_SECONDS: Final = 30
DEFAULT_RETRY_COUNT: Final = 3
DEFAULT_RETRY_INTERVAL_SECONDS: Final = 1
DEFAULT_HISTOGRAM_INTERVAL_MS: Final = 3_600_000
DEFAULT_SEARCH_WORD: Final = "#message:(UpdateMemberData) AND seqId:>1 AND namespace:ezr-main"
DEFAULT_APP_NAME: Final = "EZP.Open.Api"
DEFAULT_CLUSTER: Final = "QCloud"
DEFAULT_SORT_MODE: Final = "2"
EXCEL_MAX_CELL_LENGTH: Final = 32_767
EXCEL_HEADER_FILL: Final = "D9EAF7"
EXPORT_HEADERS: Final = (
    "品牌ID", "自然日", "日志时间", "sessionId",
    "traceId", "statusCode", "requestContent",
)

LOGGER = logging.getLogger("UpdateMemberData")


@dataclass(frozen=True)
class QueryConfig:
    brand_ids: tuple[str, ...]
    start_at: datetime
    end_at: datetime
    timezone: ZoneInfo
    token: str
    cookie: str
    page_size: int
    timeout_seconds: int
    retry_count: int
    output_path: Path
    app_name: str = DEFAULT_APP_NAME
    cluster: str = DEFAULT_CLUSTER
    search_word: str = DEFAULT_SEARCH_WORD

    def validate(self) -> None:
        if not self.brand_ids:
            raise ValueError("品牌 ID 不能为空")
        if any(not brand_id.isdigit() for brand_id in self.brand_ids):
            raise ValueError("品牌 ID 必须为数字")
        if self.start_at > self.end_at:
            raise ValueError("开始日期不能晚于结束日期")
        if not self.token.strip():
            raise ValueError("日志接口 token 不能为空")
        if self.page_size <= 0 or self.timeout_seconds <= 0 or self.retry_count <= 0:
            raise ValueError("page-size、timeout 和 retries 必须大于 0")


@dataclass(frozen=True)
class DateWindow:
    natural_date: date
    start: datetime
    end: datetime

    @property
    def start_ms(self) -> int:
        return int(self.start.timestamp() * 1000)

    @property
    def end_ms(self) -> int:
        return int(self.end.timestamp() * 1000)


@dataclass(frozen=True)
class LogRecord:
    brand_id: str
    natural_date: date
    created_on: int
    session_id: str
    trace_id: str
    status_code: str
    request_content: str

    def unique_key(self) -> tuple[str, int, str, str]:
        return self.brand_id, self.created_on, self.session_id, self.request_content


@dataclass(frozen=True)
class QuerySummary:
    brand_id: str
    natural_date: date
    api_total: int | None
    exported_count: int
    skipped_count: int


class DateWindowFactory:
    @staticmethod
    def create(start_at: datetime, end_at: datetime) -> list[DateWindow]:
        windows: list[DateWindow] = []
        current_start = start_at
        while current_start <= end_at:
            current_date = current_start.date()
            next_day_start = datetime.combine(
                current_date + timedelta(days=1), datetime_time.min, current_start.tzinfo
            )
            windows.append(
                DateWindow(
                    natural_date=current_date,
                    start=current_start,
                    end=min(next_day_start - timedelta(milliseconds=1), end_at),
                )
            )
            current_start = next_day_start
        return windows


class LogApiClient:
    def __init__(self, config: QueryConfig) -> None:
        self._config = config
        self._session = requests.Session()
        headers = {
            "Accept": "application/json, text/plain, */*",
            "Content-Type": "application/json;charset=UTF-8",
            "User-Agent": "UpdateMemberDataExporter/1.0",
        }
        self._session.headers.update(headers)
        self._add_raw_cookies(config.cookie)
        self._session.cookies.set(
            LOG_API_AUTH_COOKIE_NAME,
            json.dumps({"token": config.token}, ensure_ascii=False, separators=(",", ":")),
            domain=LOG_API_COOKIE_DOMAIN,
            path=LOG_API_COOKIE_PATH,
        )

    def _add_raw_cookies(self, raw_cookie: str) -> None:
        for cookie_part in raw_cookie.split(";"):
            normalized_part = cookie_part.strip()
            if not normalized_part:
                continue
            if "=" not in normalized_part:
                LOGGER.warning("忽略格式错误的 Cookie 片段：%s", normalized_part)
                continue
            name, value = normalized_part.split("=", maxsplit=1)
            if not name.strip():
                LOGGER.warning("忽略名称为空的 Cookie")
                continue
            self._session.cookies.set(
                name.strip(),
                value.strip(),
                domain=LOG_API_COOKIE_DOMAIN,
                path=LOG_API_COOKIE_PATH,
            )

    def close(self) -> None:
        self._session.close()

    def fetch(self, brand_id: str, window: DateWindow) -> tuple[list[LogRecord], QuerySummary]:
        page = 0
        records: list[LogRecord] = []
        seen_keys: set[tuple[str, int, str, str]] = set()
        skipped_count = 0
        api_total: int | None = None
        previous_page_signature: tuple[str, ...] | None = None

        while True:
            response_data = self._post_json(self._build_payload(brand_id, window, page))
            data_node = response_data.get("data")
            if not isinstance(data_node, dict):
                raise ValueError("接口响应缺少 data 对象")
            if api_total is None:
                api_total = self._to_optional_int(data_node.get("total"))
            events_node = data_node.get("logEvents") or []
            if not isinstance(events_node, list):
                raise ValueError("接口响应 data.logEvents 不是数组")
            LOGGER.info(
                "品牌 %s，日期 %s，第 %s 页返回 %s 条，接口总数 %s",
                brand_id, window.natural_date, page, len(events_node),
                api_total if api_total is not None else "未知",
            )
            if not events_node:
                break

            page_signature = tuple(
                self._build_event_signature(event_node) for event_node in events_node
            )
            if page_signature == previous_page_signature:
                raise RuntimeError(
                    f"接口连续返回相同分页，已停止查询以避免死循环：品牌={brand_id}，"
                    f"日期={window.natural_date}，页码={page}"
                )
            previous_page_signature = page_signature

            for event_node in events_node:
                record = self._parse_record(event_node, brand_id, window.natural_date)
                if record is None:
                    skipped_count += 1
                    continue
                unique_key = record.unique_key()
                if unique_key in seen_keys:
                    LOGGER.warning("跳过重复日志：sessionId=%s", record.session_id)
                    skipped_count += 1
                    continue
                seen_keys.add(unique_key)
                records.append(record)
            page += 1

        return records, QuerySummary(
            brand_id, window.natural_date, api_total, len(records), skipped_count
        )

    def _build_payload(self, brand_id: str, window: DateWindow, page: int) -> dict[str, object]:
        return {
            "appName": self._config.app_name,
            "cluster": self._config.cluster,
            "clusters": self._config.cluster,
            "hosts": "",
            "levels": "",
            "logCategories": "",
            "brandId": brand_id,
            "startDate": window.start_ms,
            "endDate": window.end_ms,
            "searchWord": self._config.search_word,
            "highlight": True,
            "histogram": True,
            "interval": DEFAULT_HISTOGRAM_INTERVAL_MS,
            "sortMode": DEFAULT_SORT_MODE,
            "page": page,
            "pageSize": self._config.page_size,
            "token": self._config.token,
        }

    def _post_json(self, payload: dict[str, object]) -> dict[str, object]:
        last_error: Exception | None = None
        for attempt in range(1, self._config.retry_count + 1):
            try:
                response = self._session.post(
                    LOG_API_URL, json=payload, timeout=self._config.timeout_seconds
                )
                response.raise_for_status()
                try:
                    response_data = response.json()
                except requests.JSONDecodeError as exc:
                    content_type = response.headers.get("Content-Type", "未知")
                    response_preview = response.text[:300].replace("\r", " ").replace("\n", " ")
                    raise ValueError(
                        f"日志接口返回非 JSON 内容，Content-Type={content_type}，"
                        f"响应片段={response_preview!r}"
                    ) from exc
                if not isinstance(response_data, dict):
                    raise ValueError("日志接口响应不是 JSON 对象")
                if response_data.get("status") is False:
                    raise RuntimeError(
                        f"日志接口返回失败：{response_data.get('message') or '未知错误'}"
                    )
                return response_data
            except (requests.RequestException, ValueError, RuntimeError) as exc:
                last_error = exc
                LOGGER.warning(
                    "接口请求失败，第 %s/%s 次：%s",
                    attempt, self._config.retry_count, exc,
                )
                if attempt < self._config.retry_count:
                    time.sleep(DEFAULT_RETRY_INTERVAL_SECONDS * attempt)
        raise RuntimeError(f"日志接口请求最终失败：{last_error}") from last_error

    @staticmethod
    def _parse_record(event_node: object, brand_id: str, natural_date: date) -> LogRecord | None:
        if not isinstance(event_node, dict):
            LOGGER.warning("跳过非对象类型的日志事件")
            return None
        request_content_node = event_node.get("requestContent")
        if not isinstance(request_content_node, str) or not request_content_node.strip():
            LOGGER.warning(
                "跳过缺少 requestContent 的日志：品牌 %s，日期 %s，sessionId=%s",
                brand_id, natural_date, event_node.get("sessionId") or "",
            )
            return None
        if len(request_content_node) > EXCEL_MAX_CELL_LENGTH:
            raise ValueError(
                "requestContent 超过 Excel 单元格长度上限，"
                f"品牌={brand_id}，sessionId={event_node.get('sessionId') or ''}"
            )
        created_on = LogApiClient._to_optional_int(event_node.get("createdOn")) or 0
        return LogRecord(
            brand_id=brand_id,
            natural_date=natural_date,
            created_on=created_on,
            session_id=str(event_node.get("sessionId") or ""),
            trace_id=str(event_node.get("traceId") or ""),
            status_code=str(event_node.get("statusCode") or ""),
            request_content=request_content_node,
        )

    @staticmethod
    def _to_optional_int(value: object) -> int | None:
        if value is None or isinstance(value, bool):
            return None
        try:
            return int(value)
        except (TypeError, ValueError):
            return None

    @staticmethod
    def _build_event_signature(event_node: object) -> str:
        if not isinstance(event_node, dict):
            return repr(event_node)
        return "|".join(
            (
                str(event_node.get("id") or ""),
                str(event_node.get("sessionId") or ""),
                str(event_node.get("createdOn") or ""),
                str(event_node.get("requestContent") or ""),
            )
        )


class ExcelExporter:
    def __init__(self, timezone: ZoneInfo) -> None:
        self._timezone = timezone

    def export(
        self, output_path: Path, records: list[LogRecord], summaries: list[QuerySummary]
    ) -> None:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        data_sheet = workbook.active
        data_sheet.title = "请求数据"
        data_sheet.append(EXPORT_HEADERS)
        for record in records:
            data_sheet.append(
                (
                    record.brand_id,
                    record.natural_date.isoformat(),
                    self._format_created_on(record.created_on),
                    record.session_id,
                    record.trace_id,
                    record.status_code,
                    record.request_content,
                )
            )
        self._format_sheet(data_sheet, (12, 14, 24, 36, 36, 12, 100))
        for row in data_sheet.iter_rows(min_row=2):
            row[-1].alignment = Alignment(wrap_text=True, vertical="top")

        summary_sheet = workbook.create_sheet("执行摘要")
        summary_sheet.append(("品牌ID", "自然日", "接口总数", "导出数", "跳过数"))
        for summary in summaries:
            summary_sheet.append(
                (
                    summary.brand_id,
                    summary.natural_date.isoformat(),
                    summary.api_total if summary.api_total is not None else "",
                    summary.exported_count,
                    summary.skipped_count,
                )
            )
        self._format_sheet(summary_sheet, (12, 14, 14, 12, 12))

        temporary_path = output_path.with_name(f"{output_path.stem}.tmp{output_path.suffix}")
        try:
            workbook.save(temporary_path)
            temporary_path.replace(output_path)
        finally:
            workbook.close()
            if temporary_path.exists():
                temporary_path.unlink()

    def _format_created_on(self, created_on: int) -> str:
        if created_on <= 0:
            return ""
        try:
            return datetime.fromtimestamp(created_on / 1000, self._timezone).strftime(
                "%Y-%m-%d %H:%M:%S.%f"
            )[:-3]
        except (OverflowError, OSError, ValueError):
            LOGGER.warning("createdOn 时间戳无效：%s", created_on)
            return str(created_on)

    @staticmethod
    def _format_sheet(sheet: object, widths: tuple[int, ...]) -> None:
        # openpyxl 的 Worksheet 没有稳定的公开类型入口，此处接收实际工作表对象。
        for cell in sheet[1]:  # type: ignore[index]
            cell.font = Font(bold=True)
            cell.fill = PatternFill(fill_type="solid", fgColor=EXCEL_HEADER_FILL)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.freeze_panes = "A2"  # type: ignore[attr-defined]
        sheet.auto_filter.ref = sheet.dimensions  # type: ignore[attr-defined]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width  # type: ignore[attr-defined]


class MemberDataExportApplication:
    def __init__(self, config: QueryConfig) -> None:
        self._config = config

    def run(self) -> int:
        self._config.validate()
        windows = DateWindowFactory.create(
            self._config.start_at, self._config.end_at
        )
        all_records: list[LogRecord] = []
        summaries: list[QuerySummary] = []
        client = LogApiClient(self._config)
        try:
            for brand_id in self._config.brand_ids:
                for window in windows:
                    LOGGER.info("开始查询品牌 %s，日期 %s", brand_id, window.natural_date)
                    records, summary = client.fetch(brand_id, window)
                    all_records.extend(records)
                    summaries.append(summary)
        finally:
            client.close()
        ExcelExporter(self._config.timezone).export(
            self._config.output_path, all_records, summaries
        )
        LOGGER.info("导出完成：共 %s 条数据，文件 %s", len(all_records), self._config.output_path)
        return 0


def parse_datetime_value(value: str, *, end_of_day_for_date: bool) -> datetime:
    normalized_value = value.strip()
    formats = (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d %H",
        "%Y-%m-%d",
    )
    for date_format in formats:
        try:
            parsed = datetime.strptime(normalized_value, date_format)
            if date_format == "%Y-%m-%d" and end_of_day_for_date:
                return parsed.replace(hour=23, minute=59, second=59, microsecond=999000)
            return parsed
        except ValueError:
            continue
    raise argparse.ArgumentTypeError(
        "时间格式必须为 YYYY-MM-DD、YYYY-MM-DD HH、"
        "YYYY-MM-DD HH:MM 或 YYYY-MM-DD HH:MM:SS"
    )


def parse_start_datetime(value: str) -> datetime:
    return parse_datetime_value(value, end_of_day_for_date=False)


def parse_end_datetime(value: str) -> datetime:
    return parse_datetime_value(value, end_of_day_for_date=True)


def parse_brand_ids(value: str) -> tuple[str, ...]:
    brand_ids = tuple(dict.fromkeys(item.strip() for item in value.split(",") if item.strip()))
    if not brand_ids:
        raise argparse.ArgumentTypeError("至少需要一个品牌 ID")
    if any(not brand_id.isdigit() for brand_id in brand_ids):
        raise argparse.ArgumentTypeError("品牌 ID 必须为数字，多个品牌使用英文逗号分隔")
    return brand_ids


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="按品牌和自然日查询 UpdateMemberData 日志，并导出 requestContent 到 Excel"
    )
    parser.add_argument("--brands", type=parse_brand_ids, default=DEFAULT_BRAND_IDS)
    parser.add_argument(
        "--start-date",
        type=parse_start_datetime,
        default=parse_start_datetime(DEFAULT_START_DATE),
        help="开始时间，支持日期、小时、分钟或秒",
    )
    parser.add_argument(
        "--end-date",
        type=parse_end_datetime,
        default=parse_end_datetime(DEFAULT_END_DATE),
        help="结束时间；仅输入日期时包含该自然日全天",
    )
    parser.add_argument("--token", default=DEFAULT_TOKEN, help="日志接口 token")
    parser.add_argument("--cookie", default=DEFAULT_COOKIE, help="日志平台 Cookie")
    parser.add_argument("--timezone", default=DEFAULT_TIMEZONE, help="IANA 时区名称")
    parser.add_argument("--page-size", type=int, default=DEFAULT_PAGE_SIZE)
    parser.add_argument("--timeout", type=int, default=DEFAULT_TIMEOUT_SECONDS)
    parser.add_argument("--retries", type=int, default=DEFAULT_RETRY_COUNT)
    parser.add_argument("--search-word", default=DEFAULT_SEARCH_WORD)
    parser.add_argument("--app-name", default=DEFAULT_APP_NAME)
    parser.add_argument("--cluster", default=DEFAULT_CLUSTER)
    parser.add_argument("--output", default="", help="Excel 输出路径")
    return parser.parse_args()


def create_config(args: argparse.Namespace) -> QueryConfig:
    try:
        timezone = ZoneInfo(args.timezone)
    except ZoneInfoNotFoundError as exc:
        raise ValueError(f"无效时区：{args.timezone}") from exc
    start_at = args.start_date.replace(tzinfo=timezone)
    end_at = args.end_date.replace(tzinfo=timezone)
    output_path = (
        Path(args.output).expanduser().resolve()
        if args.output.strip()
        else Path(__file__).resolve().parent
        / (
            f"UpdateMemberData_{start_at:%Y-%m-%d_%H%M%S}_"
            f"{end_at:%Y-%m-%d_%H%M%S}.xlsx"
        )
    )
    return QueryConfig(
        brand_ids=args.brands,
        start_at=start_at,
        end_at=end_at,
        timezone=timezone,
        token=args.token,
        cookie=args.cookie,
        page_size=args.page_size,
        timeout_seconds=args.timeout,
        retry_count=args.retries,
        output_path=output_path,
        app_name=args.app_name,
        cluster=args.cluster,
        search_word=args.search_word,
    )


def configure_logging() -> None:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )


def main() -> int:
    configure_logging()
    try:
        return MemberDataExportApplication(create_config(parse_args())).run()
    except KeyboardInterrupt:
        LOGGER.warning("用户中断执行")
        return 130
    except Exception as exc:
        LOGGER.exception("执行失败：%s", exc)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())

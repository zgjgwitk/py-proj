#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import logging
import shutil
import sys
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Final

import requests
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

# python "d:\Github\py-proj\京东会员通补数据-临时(8.25)\补推数据.py" --execute
PUSH_URL: Final = "http://open-tp.ezrpro.com/api/JdVipMatchVersion/UpdateMemberData"
DEFAULT_INPUT_FILE: Final = "UpdateMemberData_357261_jd.xlsx"
DEFAULT_OUTPUT_SUFFIX: Final = "_补推结果.xlsx"
DEFAULT_TIMEOUT_SECONDS: Final = 30
DEFAULT_CHECKPOINT_INTERVAL: Final = 20
DEFAULT_DELAY_MILLISECONDS: Final = 100
EXCEL_MAX_CELL_LENGTH: Final = 32_767
RESPONSE_PREVIEW_LENGTH: Final = 2_000
HEADER_FILL_COLOR: Final = "FCE4D6"
REQUEST_CONTENT_HEADER: Final = "requestContent"
BRAND_ID_HEADER: Final = "品牌ID"
RESULT_HEADERS: Final = ("补推状态", "HTTP状态码", "响应内容", "补推时间")
STATUS_SUCCESS: Final = "成功"
STATUS_FAILED: Final = "失败"
STATUS_UNKNOWN: Final = "结果未知"
STATUS_SKIPPED: Final = "跳过"
CONTENT_TYPE: Final = "application/x-www-form-urlencoded; charset=utf-8"
EZR_ENV_TAG_HEADER: Final = "ezr-env-tag"
EZR_ENV_TAG: Final = "p23414"
EZR_BRAND_ID_HEADER: Final = "ezr-brand-id"

LOGGER = logging.getLogger("MemberDataRepush")


@dataclass(frozen=True)
class PushConfig:
    input_path: Path
    output_path: Path
    execute: bool
    retry_failed: bool
    limit: int
    timeout_seconds: int
    checkpoint_interval: int
    delay_milliseconds: int

    def validate(self) -> None:
        if not self.input_path.is_file():
            raise FileNotFoundError(f"输入文件不存在：{self.input_path}")
        if self.input_path.resolve() == self.output_path.resolve():
            raise ValueError("输出文件不能覆盖输入文件，请指定独立结果文件")
        if self.limit < 0:
            raise ValueError("limit 不能小于 0")
        if self.timeout_seconds <= 0:
            raise ValueError("timeout 必须大于 0")
        if self.checkpoint_interval <= 0:
            raise ValueError("checkpoint-interval 必须大于 0")
        if self.delay_milliseconds < 0:
            raise ValueError("delay-ms 不能小于 0")


@dataclass(frozen=True)
class ColumnIndexes:
    brand_id: int
    request_content: int
    status: int
    http_status: int
    response: int
    pushed_at: int


@dataclass(frozen=True)
class PushResult:
    status: str
    http_status: int | None
    response_text: str


@dataclass
class ExecutionSummary:
    total_rows: int = 0
    valid_rows: int = 0
    processed: int = 0
    success: int = 0
    failed: int = 0
    unknown: int = 0
    skipped: int = 0


class WorkbookRepository:
    def __init__(self, config: PushConfig) -> None:
        self._config = config
        self._workbook = None
        self._sheet: Worksheet | None = None
        self._columns: ColumnIndexes | None = None

    @property
    def sheet(self) -> Worksheet:
        if self._sheet is None:
            raise RuntimeError("工作簿尚未打开")
        return self._sheet

    @property
    def columns(self) -> ColumnIndexes:
        if self._columns is None:
            raise RuntimeError("尚未解析工作表列")
        return self._columns

    def open_for_validation(self) -> None:
        self._workbook = load_workbook(self._config.input_path, read_only=True, data_only=True)
        self._sheet = self._workbook.worksheets[0]
        self._columns = self._find_columns(self._sheet, create_result_columns=False)

    def open_for_execution(self) -> None:
        if not self._config.output_path.exists():
            self._config.output_path.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(self._config.input_path, self._config.output_path)
            LOGGER.info("已创建结果文件：%s", self._config.output_path)
        self._workbook = load_workbook(self._config.output_path)
        self._sheet = self._workbook.worksheets[0]
        self._columns = self._find_columns(self._sheet, create_result_columns=True)

    def save(self) -> None:
        if self._workbook is None:
            raise RuntimeError("工作簿尚未打开")
        temporary_path = self._config.output_path.with_name(
            f"{self._config.output_path.stem}.tmp{self._config.output_path.suffix}"
        )
        try:
            self._workbook.save(temporary_path)
            temporary_path.replace(self._config.output_path)
        finally:
            if temporary_path.exists():
                temporary_path.unlink()

    def close(self) -> None:
        if self._workbook is not None:
            self._workbook.close()

    @staticmethod
    def get_text(sheet: Worksheet, row_index: int, column_index: int) -> str:
        value = sheet.cell(row=row_index, column=column_index).value
        return str(value).strip() if value is not None else ""

    def write_result(self, row_index: int, result: PushResult) -> None:
        pushed_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        values = (
            result.status,
            result.http_status if result.http_status is not None else "",
            self._safe_excel_text(result.response_text),
            pushed_at,
        )
        indexes = (
            self.columns.status,
            self.columns.http_status,
            self.columns.response,
            self.columns.pushed_at,
        )
        for column_index, value in zip(indexes, values, strict=True):
            cell = self.sheet.cell(row=row_index, column=column_index)
            cell.value = value
            if isinstance(value, str):
                # 防止服务端响应以“=”开头时被 Excel 当作公式执行。
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=column_index == self.columns.response)

    @staticmethod
    def _find_columns(sheet: Worksheet, create_result_columns: bool) -> ColumnIndexes:
        header_map: dict[str, int] = {}
        for cell in sheet[1]:
            if cell.value is not None:
                header_map[str(cell.value).strip()] = cell.column
        request_content_index = header_map.get(REQUEST_CONTENT_HEADER)
        if request_content_index is None:
            raise ValueError(f"Excel 缺少必需列：{REQUEST_CONTENT_HEADER}")
        brand_id_index = header_map.get(BRAND_ID_HEADER)
        if brand_id_index is None:
            raise ValueError(f"Excel 缺少必需列：{BRAND_ID_HEADER}")

        result_indexes: list[int] = []
        for header in RESULT_HEADERS:
            column_index = header_map.get(header)
            if column_index is None:
                if not create_result_columns:
                    result_indexes.append(0)
                    continue
                column_index = sheet.max_column + 1
                cell = sheet.cell(row=1, column=column_index, value=header)
                cell.font = Font(bold=True)
                cell.fill = PatternFill(fill_type="solid", fgColor=HEADER_FILL_COLOR)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                header_map[header] = column_index
            result_indexes.append(column_index)

        return ColumnIndexes(brand_id_index, request_content_index, *result_indexes)

    @staticmethod
    def _safe_excel_text(value: str) -> str:
        if len(value) <= EXCEL_MAX_CELL_LENGTH:
            return value
        suffix = "\n[响应内容超过 Excel 单元格上限，已截断]"
        return value[: EXCEL_MAX_CELL_LENGTH - len(suffix)] + suffix


class MemberDataPushClient:
    def __init__(self, timeout_seconds: int) -> None:
        self._timeout_seconds = timeout_seconds
        self._session = requests.Session()
        self._session.headers.update(
            {
                "Accept": "application/json, text/plain, */*",
                "Content-Type": CONTENT_TYPE,
                # EZR_ENV_TAG_HEADER: EZR_ENV_TAG,
                "User-Agent": "MemberDataRepush/1.0",
            }
        )

    def close(self) -> None:
        self._session.close()

    def push(self, request_content: str, brand_id: str) -> PushResult:
        request_url = f"{PUSH_URL}?{request_content}"
        try:
            response = self._session.post(
                request_url,
                data=b"",
                headers={EZR_BRAND_ID_HEADER: brand_id},
                timeout=self._timeout_seconds,
                allow_redirects=False,
            )
        except requests.RequestException as exc:
            # POST 超时或连接中断时无法确认服务端是否已经处理，禁止自动重试。
            return PushResult(STATUS_UNKNOWN, None, f"请求异常：{exc}")

        response_text = response.text
        if response.is_redirect or response.is_permanent_redirect:
            location = response.headers.get("Location", "")
            return PushResult(
                STATUS_FAILED,
                response.status_code,
                f"接口返回重定向，Location={location}",
            )
        if not 200 <= response.status_code < 300:
            return PushResult(STATUS_FAILED, response.status_code, response_text)
        if not self._is_business_success(response_text):
            return PushResult(STATUS_FAILED, response.status_code, response_text)
        return PushResult(STATUS_SUCCESS, response.status_code, response_text)

    @staticmethod
    def _is_business_success(response_text: str) -> bool:
        try:
            response_data = json.loads(response_text)
        except json.JSONDecodeError:
            return True
        if not isinstance(response_data, dict):
            return True
        return response_data.get("status") is not False


class MemberDataPushApplication:
    def __init__(self, config: PushConfig) -> None:
        self._config = config

    def run(self) -> int:
        self._config.validate()
        if not self._config.execute:
            return self._validate_only()
        return self._execute()

    def _validate_only(self) -> int:
        repository = WorkbookRepository(self._config)
        try:
            repository.open_for_validation()
            summary = self._inspect_rows(repository)
        finally:
            repository.close()
        LOGGER.info(
            "校验完成：总行数 %s，有效补推数据 %s，无效或空值 %s",
            summary.total_rows,
            summary.valid_rows,
            summary.total_rows - summary.valid_rows,
        )
        LOGGER.info("当前为校验模式，未发送请求；确认后请增加 --execute")
        return 0

    def _execute(self) -> int:
        repository = WorkbookRepository(self._config)
        client = MemberDataPushClient(self._config.timeout_seconds)
        summary = ExecutionSummary()
        unsaved_count = 0
        try:
            repository.open_for_execution()
            summary.total_rows = max(repository.sheet.max_row - 1, 0)
            for row_index in range(2, repository.sheet.max_row + 1):
                if self._config.limit and summary.processed >= self._config.limit:
                    LOGGER.info("达到 limit=%s，停止发送新请求", self._config.limit)
                    break
                request_content = repository.get_text(
                    repository.sheet, row_index, repository.columns.request_content
                )
                brand_id = repository.get_text(
                    repository.sheet, row_index, repository.columns.brand_id
                )
                existing_status = repository.get_text(
                    repository.sheet, row_index, repository.columns.status
                )
                if not request_content:
                    summary.skipped += 1
                    if not existing_status:
                        repository.write_result(
                            row_index, PushResult(STATUS_SKIPPED, None, "requestContent 为空")
                        )
                        unsaved_count += 1
                    continue
                if not self._is_valid_brand_id(brand_id):
                    summary.skipped += 1
                    if not existing_status:
                        repository.write_result(
                            row_index,
                            PushResult(STATUS_SKIPPED, None, "品牌ID为空或不是数字"),
                        )
                        unsaved_count += 1
                    continue
                summary.valid_rows += 1
                if existing_status and not self._should_retry(existing_status):
                    summary.skipped += 1
                    continue

                result = client.push(request_content, brand_id)
                repository.write_result(row_index, result)
                summary.processed += 1
                unsaved_count += 1
                self._record_result(summary, result)
                LOGGER.info(
                    "[%s/%s] 第 %s 行补推%s，HTTP=%s",
                    summary.processed,
                    self._config.limit or summary.total_rows,
                    row_index,
                    result.status,
                    result.http_status if result.http_status is not None else "未知",
                )

                if unsaved_count >= self._config.checkpoint_interval:
                    repository.save()
                    unsaved_count = 0
                    LOGGER.info("已保存执行进度：累计处理 %s 条", summary.processed)
                if self._config.delay_milliseconds:
                    time.sleep(self._config.delay_milliseconds / 1000)
        finally:
            try:
                if unsaved_count > 0:
                    repository.save()
            finally:
                client.close()
                repository.close()

        LOGGER.info(
            "执行完成：处理 %s，成功 %s，失败 %s，结果未知 %s，跳过 %s",
            summary.processed,
            summary.success,
            summary.failed,
            summary.unknown,
            summary.skipped,
        )
        LOGGER.info("结果文件：%s", self._config.output_path)
        return 0 if summary.failed == 0 and summary.unknown == 0 else 2

    def _inspect_rows(self, repository: WorkbookRepository) -> ExecutionSummary:
        summary = ExecutionSummary(total_rows=max(repository.sheet.max_row - 1, 0))
        # 只读工作表必须顺序遍历；逐行调用 cell() 会反复扫描 XML，数据量大时极慢。
        first_column = min(repository.columns.brand_id, repository.columns.request_content)
        last_column = max(repository.columns.brand_id, repository.columns.request_content)
        brand_offset = repository.columns.brand_id - first_column
        request_offset = repository.columns.request_content - first_column
        for row_values in repository.sheet.iter_rows(
            min_row=2,
            min_col=first_column,
            max_col=last_column,
            values_only=True,
        ):
            brand_id = str(row_values[brand_offset] or "").strip()
            request_content = str(row_values[request_offset] or "").strip()
            if request_content and self._is_valid_brand_id(brand_id):
                summary.valid_rows += 1
        return summary

    @staticmethod
    def _is_valid_brand_id(brand_id: str) -> bool:
        return bool(brand_id) and brand_id.isdigit()

    def _should_retry(self, existing_status: str) -> bool:
        return self._config.retry_failed and existing_status == STATUS_FAILED

    @staticmethod
    def _record_result(summary: ExecutionSummary, result: PushResult) -> None:
        if result.status == STATUS_SUCCESS:
            summary.success += 1
        elif result.status == STATUS_UNKNOWN:
            summary.unknown += 1
        else:
            summary.failed += 1


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="读取 UpdateMemberData Excel，并将 requestContent 补推到接口"
    )
    parser.add_argument("--input", default=DEFAULT_INPUT_FILE, help="输入 Excel 路径")
    parser.add_argument("--output", default="", help="结果 Excel 路径")
    parser.add_argument("--execute", action="store_true", help="实际发送生产补推请求")
    parser.add_argument(
        "--retry-failed",
        action="store_true",
        help="重新请求结果文件中状态为‘失败’的行，不重试‘结果未知’",
    )
    parser.add_argument("--limit", type=int, default=0, help="本次最多发送数量；0 表示不限")
    parser.add_argument("--timeout", type=int, default=DEFAULT_TIMEOUT_SECONDS)
    parser.add_argument(
        "--checkpoint-interval", type=int, default=DEFAULT_CHECKPOINT_INTERVAL
    )
    parser.add_argument("--delay-ms", type=int, default=DEFAULT_DELAY_MILLISECONDS)
    return parser.parse_args()


def create_config(args: argparse.Namespace) -> PushConfig:
    script_directory = Path(__file__).resolve().parent
    input_path = Path(args.input).expanduser()
    if not input_path.is_absolute():
        input_path = script_directory / input_path
    input_path = input_path.resolve()

    if args.output.strip():
        output_path = Path(args.output).expanduser()
        if not output_path.is_absolute():
            output_path = script_directory / output_path
        output_path = output_path.resolve()
    else:
        output_path = input_path.with_name(f"{input_path.stem}{DEFAULT_OUTPUT_SUFFIX}")

    return PushConfig(
        input_path=input_path,
        output_path=output_path,
        execute=args.execute,
        retry_failed=args.retry_failed,
        limit=args.limit,
        timeout_seconds=args.timeout,
        checkpoint_interval=args.checkpoint_interval,
        delay_milliseconds=args.delay_ms,
    )


def configure_console() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8")
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )


def main() -> int:
    configure_console()
    try:
        return MemberDataPushApplication(create_config(parse_args())).run()
    except KeyboardInterrupt:
        LOGGER.warning("用户中断执行，已保存最近进度")
        return 130
    except Exception as exc:
        LOGGER.exception("执行失败：%s", exc)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
